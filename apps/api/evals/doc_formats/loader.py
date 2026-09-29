"""Read a bake-off document of any supported format into plain text and an import upload.

The format comes from the file extension and is confirmed by the content: Office files are ZIP
archives holding their main part, a PDF carries `%PDF-` in its first kilobyte, an HTML file
holds tags, and a text file is UTF-8 that is neither a ZIP nor a PDF. A PDF whose text layer
averages fewer than 25 non-space characters per page is scanned and is read by OCR.

Run: `uv run python -m evals.doc_formats.loader <path> [--ocr-fake]`.
"""

from __future__ import annotations

import argparse
import asyncio
import io
import re
import sys
import zipfile
from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

import openpyxl
from pptx import Presentation
from pptx.shapes.group import GroupShape

from app.utilities.document_text import (
    DocumentTooLargeError,
    DocumentUnreadableError,
    _docx_paragraphs,
    _pdf_pages,
)
from evals.doc_formats.loaded_document import LoadedDocument, OcrUsage
from evals.doc_formats.ocr_client import FakeOcrClient, OcrClient
from evals.html_to_text import html_to_markdown

FORMAT_BY_SUFFIX = {
    ".txt": "txt",
    ".text": "txt",
    ".md": "md",
    ".markdown": "md",
    ".html": "html",
    ".htm": "html",
    ".docx": "docx",
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".xlsx": "xlsx",
}
DOCUMENT_SUFFIXES = frozenset(FORMAT_BY_SUFFIX)
SCANNED_PDF_CHARS_PER_PAGE = 25
PREVIEW_CHARS = 800

_ZIP_SIGNATURE = b"PK\x03\x04"
_PDF_SIGNATURE = b"%PDF-"
_PDF_SIGNATURE_WINDOW = 1024
_OFFICE_MAIN_PART = {
    "docx": "word/document.xml",
    "pptx": "ppt/presentation.xml",
    "xlsx": "xl/workbook.xml",
}
_NATIVE_UPLOAD_FORMATS = frozenset({"txt", "md", "docx", "pdf"})
_HTML_TAG = re.compile(r"<(!doctype\s+html|html|head|body|p|div|h[1-6]|section|article)\b", re.I)
_BLANK_RUNS = re.compile(r"\n{3,}")


class DocumentFormatError(Exception):
    """The document is of an unsupported format, does not match its extension, or cannot be
    read without a missing collaborator such as an OCR client."""


async def load_document(path: Path, ocr: OcrClient | None = None) -> LoadedDocument:
    """Read `path` into its full text and the upload to send to the sentence import."""
    fmt = FORMAT_BY_SUFFIX.get(path.suffix.lower())
    if fmt is None:
        raise DocumentFormatError(f"{path.name}: unsupported document extension {path.suffix!r}")
    data = path.read_bytes()
    _confirm_format(path, fmt, data)
    usage: OcrUsage | None = None
    pages: int | None = None
    try:
        if fmt in ("txt", "md"):
            text = _decode(path, data)
        elif fmt == "html":
            text = html_to_markdown(_decode(path, data))
        elif fmt == "docx":
            text = _blocks(_docx_paragraphs(data))
        elif fmt == "pptx":
            text, pages = _pptx_text(data)
        elif fmt == "xlsx":
            text = _xlsx_text(data)
        else:
            page_texts = _pdf_pages(data)
            pages = len(page_texts)
            if not _is_scanned(page_texts):
                text = _blocks(page_texts)
            else:
                if ocr is None:
                    raise DocumentFormatError(
                        f"{path.name}: the PDF is scanned (no usable text layer) and needs an "
                        "OCR client"
                    )
                fmt = "scanned_pdf"
                result = await ocr.ocr_pdf(data)
                text = _normalise(result.text)
                usage = OcrUsage(ocr.deployment, result.pages, result.latency_ms, result.cost_eur)
                pages = result.pages or pages
    except (DocumentUnreadableError, DocumentTooLargeError) as exc:
        raise DocumentFormatError(f"{path.name}: {exc}") from exc
    if fmt in _NATIVE_UPLOAD_FORMATS:
        upload_name, upload_bytes = path.name, data
    else:
        upload_name, upload_bytes = f"{path.stem}.md", text.encode("utf-8")
    return LoadedDocument(
        path=path,
        format=fmt,
        text=text,
        upload_name=upload_name,
        upload_bytes=upload_bytes,
        ocr=usage,
        pages=pages,
        words=len(text.split()),
    )


def _confirm_format(path: Path, fmt: str, data: bytes) -> None:
    is_zip = data.startswith(_ZIP_SIGNATURE)
    is_pdf = _PDF_SIGNATURE in data[:_PDF_SIGNATURE_WINDOW]
    if fmt in _OFFICE_MAIN_PART:
        member = _OFFICE_MAIN_PART[fmt]
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                present = member in archive.namelist()
        except zipfile.BadZipFile:
            present = False
        if not (is_zip and present):
            raise DocumentFormatError(f"{path.name}: not a {fmt} file (no {member} in a ZIP)")
    elif fmt == "pdf":
        if not is_pdf:
            raise DocumentFormatError(f"{path.name}: not a PDF (no %PDF- header)")
    else:
        if is_zip or is_pdf:
            raise DocumentFormatError(f"{path.name}: a ZIP or PDF file, not {fmt} text")
        if fmt == "html" and not _HTML_TAG.search(data[:65536].decode("utf-8", "replace")):
            raise DocumentFormatError(f"{path.name}: no HTML tags found")


def _decode(path: Path, data: bytes) -> str:
    try:
        return _normalise(data.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise DocumentFormatError(f"{path.name}: not UTF-8 text") from exc


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return _BLANK_RUNS.sub("\n\n", text).strip() + "\n"


def _blocks(blocks: Iterable[str]) -> str:
    return _normalise("\n\n".join(b.strip() for b in blocks if b and b.strip()))


def _is_scanned(page_texts: list[str]) -> bool:
    if not page_texts:
        return True
    chars = sum(len(re.sub(r"\s", "", t)) for t in page_texts)
    return chars / len(page_texts) < SCANNED_PDF_CHARS_PER_PAGE


def _pptx_text(data: bytes) -> tuple[str, int]:
    """Per slide: `## Slide N: <title>`, then text frames and tables in shape order, then notes."""
    presentation = Presentation(io.BytesIO(data))
    blocks: list[str] = []
    slides = 0
    for number, slide in enumerate(presentation.slides, start=1):
        slides += 1
        title_shape = slide.shapes.title
        title = _one_line(title_shape.text_frame.text) if title_shape is not None else ""
        blocks.append(f"## Slide {number}: {title}" if title else f"## Slide {number}")
        for shape in _shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                blocks.extend(_paragraphs(shape.text_frame))
            elif shape.has_table:
                rows = [_row(cell.text for cell in row.cells) for row in shape.table.rows]
                blocks.append("\n".join(r for r in rows if r))
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = list(_paragraphs(frame)) if frame is not None else []
            if notes:
                blocks.append("Notes: " + " ".join(notes))
    return _blocks(blocks), slides


def _shapes(shapes: Iterable[Any]) -> Iterator[Any]:
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _paragraphs(frame: Any) -> Iterator[str]:
    for paragraph in frame.paragraphs:
        text = _one_line(paragraph.text)
        if text:
            yield text


def _xlsx_text(data: bytes) -> str:
    """Per sheet: `## <sheet>`, then each non-empty row as its cell values joined by ` | `."""
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = [_row(_cell(v) for v in row) for row in sheet.iter_rows(values_only=True)]
            blocks.append(f"## {sheet.title}")
            lines = "\n".join(r for r in rows if r)
            if lines:
                blocks.append(lines)
    finally:
        workbook.close()
    return _blocks(blocks)


def _cell(value: Any) -> str:
    return "" if value is None else str(value)


def _row(cells: Iterable[str]) -> str:
    return " | ".join(c for c in (_one_line(x) for x in cells) if c)


def _one_line(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(prog="python -m evals.doc_formats.loader")
    parser.add_argument("path", type=Path)
    parser.add_argument(
        "--ocr-fake", action="store_true", help="read scanned PDFs with a fake OCR client"
    )
    args = parser.parse_args(argv)
    ocr = FakeOcrClient("Fake OCR text for a scanned page.") if args.ocr_fake else None
    try:
        doc = asyncio.run(load_document(args.path, ocr))
    except DocumentFormatError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    print(f"format: {doc.format}")
    print(f"words: {doc.words}")
    print(f"pages: {doc.pages if doc.pages is not None else '-'}")
    print(f"upload: {doc.upload_name} ({len(doc.upload_bytes)} bytes)")
    if doc.ocr is not None:
        print(
            f"ocr: {doc.ocr.deployment}, {doc.ocr.pages} pages, {doc.ocr.latency_ms} ms, "
            f"EUR {doc.ocr.cost_eur:.4f}"
        )
    print()
    print(doc.text[:PREVIEW_CHARS])
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
