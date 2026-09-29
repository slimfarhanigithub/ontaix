"""Writes one document's gold tree in eight ontology and spreadsheet formats, and reads each back.

The source of truth is `<name>.expected.yaml`. The same logical tree - every concept's label,
parents, birth action and accepted actions, aliases, and every relation with its actions and
inverse actions - is written as:

- `<name>.ttl`, `<name>.owl`, `<name>.jsonld`: OWL 2 in Turtle, RDF/XML and JSON-LD (one graph);
  the tree as `rdfs:subClassOf`, relations as object property restrictions
  (`owl:someValuesFrom`).
- `<name>.skos.ttl`: SKOS in Turtle; `skos:broader`, and top concepts of a scheme named after the
  company; relations as direct triples between concepts, each doubled by `skos:related`.
- `<name>.obo`: OBO 1.4; `is_a`, `relationship` lines and `[Typedef]` stanzas.
- `<name>.csv`: CSV edge list (`label,parent`); no relations.
- `<name>.xlsx`: Excel outline, one column per level (the full path) plus `Label` and `Parent`
  columns; relations on a second sheet.
- `<name>.json`: nested JSON tree, the root node labelled with the company; a `relations` array.

Actions: the first action is `oxe:birthAction` (OBO `birth_action`), the rest
`oxe:acceptedAction`; in RDF the order of accepted actions is not kept, so comparisons sort them.
The company is the root: an OWL class labelled with the company name, the SKOS concept scheme,
the OBO root term, and the top of the JSON tree. A relation verb is an object property labelled
with the first action; a verb whose accepted or inverse actions differ between relations gets a
numbered property of its own.

Run: `python gold_formats.py write <name>` then `python gold_formats.py check <name>`
(`--documents DIR`, default the parent folder; outputs go to `<documents>/gold/`).
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from rdflib import BNode, Graph, Literal, Namespace, URIRef
from rdflib.namespace import OWL, RDF, RDFS, SKOS

OXE = Namespace("https://example.org/ontaix/evals/vocab#")
FORMATS = ("ttl", "owl", "jsonld", "skos.ttl", "obo", "csv", "xlsx", "json")
_RDF_FORMATS = {"ttl": "turtle", "owl": "xml", "jsonld": "json-ld", "skos.ttl": "turtle"}
_WITH_RELATIONS = {"ttl", "owl", "jsonld", "skos.ttl", "obo", "xlsx", "json"}


@dataclass(frozen=True)
class Concept:
    label: str
    parents: tuple[str, ...]
    actions: tuple[str, ...]  # first = birth action, rest sorted
    aliases: tuple[str, ...]


@dataclass(frozen=True)
class Relation:
    source: str
    target: str
    actions: tuple[str, ...]
    inverse: tuple[str, ...]


@dataclass
class Gold:
    company: str
    concepts: dict[str, Concept]
    relations: set[Relation]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["write", "check"])
    parser.add_argument("name")
    parser.add_argument("--documents", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    gold = load_yaml(args.documents / f"{args.name}.expected.yaml")
    out = args.documents / "gold"
    if args.command == "write":
        out.mkdir(exist_ok=True)
        write_all(gold, out, args.name)
        print(f"wrote {len(FORMATS)} formats of {args.name} to {out}")
        return 0
    failures = 0
    for fmt in FORMATS:
        path = out / f"{args.name}.{fmt}"
        got = read_format(path, fmt, gold.company)
        problems = compare(gold, got, relations=fmt in _WITH_RELATIONS)
        print(
            f"{'OK' if not problems else 'FAIL'} {path.name}: {len(got.concepts)} concepts, "
            f"{len(got.relations)} relations"
        )
        for p in problems[:20]:
            print(f"  {p}")
        failures += bool(problems)
    return 1 if failures else 0


def load_yaml(path: Path) -> Gold:
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    concepts = {}
    for c in raw["expected"]["concepts"]:
        concepts[c["label"]] = Concept(
            c["label"],
            tuple(sorted(_listed(c["parent"]))),
            _actions(_listed(c["action"])),
            tuple(sorted(c.get("aliases", []))),
        )
    relations = {
        Relation(
            r["from"], r["to"], _actions(_listed(r["action"])), tuple(sorted(r.get("inverse", [])))
        )
        for r in raw["expected"].get("relations", [])
    }
    return Gold(raw["company"], concepts, relations)


def compare(want: Gold, got: Gold, relations: bool) -> list[str]:
    problems = []
    if want.company != got.company:
        problems.append(f"company {got.company!r}, want {want.company!r}")
    for label in sorted(set(want.concepts) | set(got.concepts)):
        w, g = want.concepts.get(label), got.concepts.get(label)
        if w != g:
            problems.append(f"concept {label!r}: got {g}, want {w}")
    if relations and want.relations != got.relations:
        for r in sorted(want.relations - got.relations, key=str):
            problems.append(f"missing relation {r}")
        for r in sorted(got.relations - want.relations, key=str):
            problems.append(f"extra relation {r}")
    return problems


def write_all(gold: Gold, out: Path, name: str) -> None:
    base = f"https://example.org/ontaix/evals/{name}#"
    owl = owl_graph(gold, Namespace(base))
    (out / f"{name}.ttl").write_text(owl.serialize(format="turtle"), encoding="utf-8")
    (out / f"{name}.owl").write_text(owl.serialize(format="xml"), encoding="utf-8")
    (out / f"{name}.jsonld").write_text(owl.serialize(format="json-ld", indent=2), encoding="utf-8")
    skos = skos_graph(gold, Namespace(base))
    (out / f"{name}.skos.ttl").write_text(skos.serialize(format="turtle"), encoding="utf-8")
    (out / f"{name}.obo").write_text(obo_text(gold, name), encoding="utf-8", newline="\n")
    (out / f"{name}.csv").write_text(csv_text(gold), encoding="utf-8", newline="")
    (out / f"{name}.xlsx").write_bytes(xlsx_bytes(gold))
    (out / f"{name}.json").write_text(
        json.dumps(json_tree(gold), indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )


def owl_graph(gold: Gold, ns: Namespace) -> Graph:
    g = _graph(ns)
    ontology = URIRef(str(ns).rstrip("#"))
    g.add((ontology, RDF.type, OWL.Ontology))
    g.add((ontology, RDFS.label, Literal(f"{gold.company} gold tree", lang="en")))
    iri = _iris(gold, ns)
    for label, node in iri.items():
        g.add((node, RDF.type, OWL.Class))
        g.add((node, RDFS.label, Literal(label, lang="en")))
    for c in gold.concepts.values():
        for p in c.parents:
            g.add((iri[c.label], RDFS.subClassOf, iri[p]))
        _add_actions(g, iri[c.label], c.actions)
        for alias in c.aliases:
            g.add((iri[c.label], SKOS.altLabel, Literal(alias, lang="en")))
    for r, prop in _properties(gold, ns, g).items():
        restriction = BNode()
        g.add((restriction, RDF.type, OWL.Restriction))
        g.add((restriction, OWL.onProperty, prop))
        g.add((restriction, OWL.someValuesFrom, iri[r.target]))
        g.add((iri[r.source], RDFS.subClassOf, restriction))
    return g


def skos_graph(gold: Gold, ns: Namespace) -> Graph:
    g = _graph(ns)
    scheme = URIRef(str(ns).rstrip("#"))
    g.add((scheme, RDF.type, SKOS.ConceptScheme))
    g.add((scheme, SKOS.prefLabel, Literal(gold.company, lang="en")))
    iri = _iris(gold, ns)
    for c in gold.concepts.values():
        node = iri[c.label]
        g.add((node, RDF.type, SKOS.Concept))
        g.add((node, SKOS.prefLabel, Literal(c.label, lang="en")))
        g.add((node, SKOS.inScheme, scheme))
        for p in c.parents:
            if p == gold.company:
                g.add((node, SKOS.topConceptOf, scheme))
                g.add((scheme, SKOS.hasTopConcept, node))
            else:
                g.add((node, SKOS.broader, iri[p]))
        _add_actions(g, node, c.actions)
        for alias in c.aliases:
            g.add((node, SKOS.altLabel, Literal(alias, lang="en")))
    for r, prop in _properties(gold, ns, g).items():
        g.add((iri[r.source], prop, iri[r.target]))
        g.add((iri[r.source], SKOS.related, iri[r.target]))
    return g


def obo_text(gold: Gold, name: str) -> str:
    ids = {gold.company: "GOLD:0000000"}
    for i, label in enumerate(gold.concepts, start=1):
        ids[label] = f"GOLD:{i:07d}"
    typedefs = _verb_ids(gold)
    out = [
        "format-version: 1.4",
        f"ontology: {name}",
        f"remark: Gold tree of {name}.md. The company is the root term GOLD:0000000.",
        "",
        "[Term]",
        f"id: {ids[gold.company]}",
        f"name: {gold.company}",
        "",
    ]
    for c in gold.concepts.values():
        out += ["[Term]", f"id: {ids[c.label]}", f"name: {c.label}"]
        out += [f'synonym: "{_obo_quote(a)}" EXACT []' for a in c.aliases]
        out += [f"is_a: {ids[p]} ! {p}" for p in c.parents]
        out.append(f'property_value: birth_action "{_obo_quote(c.actions[0])}" xsd:string')
        out += [
            f'property_value: accepted_action "{_obo_quote(a)}" xsd:string' for a in c.actions[1:]
        ]
        for r in sorted(gold.relations, key=str):
            if r.source == c.label:
                out.append(f"relationship: {typedefs[r]} {ids[r.target]} ! {r.target}")
        out.append("")
    written: set[str] = set()
    for r, tid in sorted(typedefs.items(), key=lambda x: x[1]):
        if tid in written:
            continue
        written.add(tid)
        out += ["[Typedef]", f"id: {tid}", f"name: {r.actions[0]}"]
        out += [
            f'property_value: accepted_action "{_obo_quote(a)}" xsd:string' for a in r.actions[1:]
        ]
        out += [f'property_value: inverse_action "{_obo_quote(a)}" xsd:string' for a in r.inverse]
        out.append("")
    return "\n".join(out)


def csv_text(gold: Gold) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(["label", "parent", "action", "accepted_actions", "aliases"])
    for c in gold.concepts.values():
        for p in c.parents:
            writer.writerow(
                [c.label, p, c.actions[0], "; ".join(c.actions[1:]), "; ".join(c.aliases)]
            )
    return buffer.getvalue()


def xlsx_bytes(gold: Gold) -> bytes:
    depth = max(len(_path(gold, c.label)) for c in gold.concepts.values())
    book = Workbook()
    sheet = book.active
    sheet.title = "Hierarchy"
    header = [f"Level {i}" for i in range(depth + 1)]
    header[0] = "Company"
    sheet.append([*header, "Label", "Parent", "Action", "Accepted actions", "Aliases"])
    for c in gold.concepts.values():
        path = _path(gold, c.label)
        row = [*path, *([None] * (depth + 1 - len(path)))]
        for p in c.parents:
            sheet.append(
                [*row, c.label, p, c.actions[0], "; ".join(c.actions[1:]), "; ".join(c.aliases)]
            )
    relations = book.create_sheet("Relations")
    relations.append(["From", "Action", "Accepted actions", "Inverse actions", "To"])
    for r in sorted(gold.relations, key=str):
        relations.append(
            [r.source, r.actions[0], "; ".join(r.actions[1:]), "; ".join(r.inverse), r.target]
        )
    for ws in (sheet, relations):
        for cell in ws[1]:
            cell.font = Font(bold=True)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def json_tree(gold: Gold) -> dict:
    children: dict[str, list[str]] = {}
    for c in gold.concepts.values():
        children.setdefault(c.parents[0], []).append(c.label)

    def node(label: str) -> dict:
        c = gold.concepts[label]
        out: dict = {"label": label, "action": list(c.actions)}
        if c.aliases:
            out["aliases"] = list(c.aliases)
        if len(c.parents) > 1:
            out["alsoUnder"] = list(c.parents[1:])
        kids = children.get(label, [])
        if kids:
            out["children"] = [node(k) for k in kids]
        return out

    return {
        "label": gold.company,
        "children": [node(k) for k in children.get(gold.company, [])],
        "relations": [
            {
                "from": r.source,
                "to": r.target,
                "actions": list(r.actions),
                "inverse": list(r.inverse),
            }
            for r in sorted(gold.relations, key=str)
        ],
    }


def read_format(path: Path, fmt: str, company: str) -> Gold:
    if fmt in ("ttl", "owl", "jsonld"):
        return _read_owl(path, _RDF_FORMATS[fmt])
    if fmt == "skos.ttl":
        return _read_skos(path)
    if fmt == "obo":
        return _read_obo(path)
    if fmt == "csv":
        return _read_csv(path, company)
    if fmt == "xlsx":
        return _read_xlsx(path)
    return _read_json(path)


def _read_owl(path: Path, rdf_format: str) -> Gold:
    g = Graph().parse(path, format=rdf_format)
    ontology = next(g.subjects(RDF.type, OWL.Ontology))
    company = str(g.value(ontology, RDFS.label)).removesuffix(" gold tree")
    labels = {s: str(o) for s, o in g.subject_objects(RDFS.label) if (s, RDF.type, OWL.Class) in g}
    concepts, relations = {}, set()
    for node, label in labels.items():
        if label == company:
            continue
        parents = []
        for sup in g.objects(node, RDFS.subClassOf):
            if isinstance(sup, BNode):
                prop = g.value(sup, OWL.onProperty)
                target = labels[g.value(sup, OWL.someValuesFrom)]
                relations.add(_relation_of(g, prop, label, target))
            else:
                parents.append(labels[sup])
        concepts[label] = _concept_of(g, node, label, parents)
    return Gold(company, concepts, relations)


def _read_skos(path: Path) -> Gold:
    g = Graph().parse(path, format="turtle")
    scheme = next(g.subjects(RDF.type, SKOS.ConceptScheme))
    company = str(g.value(scheme, SKOS.prefLabel))
    labels = {s: str(g.value(s, SKOS.prefLabel)) for s in g.subjects(RDF.type, SKOS.Concept)}
    concepts, relations = {}, set()
    for node, label in labels.items():
        parents = [labels[b] for b in g.objects(node, SKOS.broader)]
        if (node, SKOS.topConceptOf, scheme) in g:
            parents.append(company)
        concepts[label] = _concept_of(g, node, label, parents)
    for s, p, o in g:
        if (p, RDF.type, OWL.ObjectProperty) in g:
            relations.add(_relation_of(g, p, labels[s], labels[o]))
    return Gold(company, concepts, relations)


def _read_obo(path: Path) -> Gold:
    stanzas = re.split(r"\n\s*\n", path.read_text(encoding="utf-8"))
    names: dict[str, str] = {}
    terms: list[dict[str, list[str]]] = []
    typedefs: dict[str, dict[str, list[str]]] = {}
    for stanza in stanzas:
        lines = [line for line in stanza.strip().splitlines() if line]
        if not lines or lines[0] not in ("[Term]", "[Typedef]"):
            continue
        tags: dict[str, list[str]] = {}
        for line in lines[1:]:
            key, _, value = line.partition(": ")
            tags.setdefault(key, []).append(value)
        if lines[0] == "[Term]":
            names[tags["id"][0]] = tags["name"][0]
            terms.append(tags)
        else:
            typedefs[tags["id"][0]] = tags
    root = names["GOLD:0000000"]
    concepts, relations = {}, set()
    for t in terms:
        label = t["name"][0]
        if label == root:
            continue
        values = _obo_values(t.get("property_value", []))
        parents = [names[v.split(" ! ")[0]] for v in t.get("is_a", [])]
        aliases = [re.match(r'"(.*)" EXACT', s).group(1) for s in t.get("synonym", [])]
        concepts[label] = Concept(
            label,
            tuple(sorted(parents)),
            _actions(values.get("birth_action", []) + values.get("accepted_action", [])),
            tuple(sorted(aliases)),
        )
        for rel in t.get("relationship", []):
            tid, target = rel.split(" ! ")[0].split(" ")
            td = typedefs[tid]
            tv = _obo_values(td.get("property_value", []))
            relations.add(
                Relation(
                    label,
                    names[target],
                    _actions(td["name"] + tv.get("accepted_action", [])),
                    tuple(sorted(tv.get("inverse_action", []))),
                )
            )
    return Gold(root, concepts, relations)


def _read_csv(path: Path, company: str) -> Gold:
    rows = list(csv.DictReader(path.read_text(encoding="utf-8").splitlines()))
    parents: dict[str, list[str]] = {}
    first: dict[str, dict[str, str]] = {}
    for row in rows:
        parents.setdefault(row["label"], []).append(row["parent"])
        first.setdefault(row["label"], row)
    concepts = {
        label: Concept(
            label,
            tuple(sorted(ps)),
            _actions([first[label]["action"], *_split(first[label]["accepted_actions"])]),
            tuple(sorted(_split(first[label]["aliases"]))),
        )
        for label, ps in parents.items()
    }
    return Gold(company, concepts, set())


def _read_xlsx(path: Path) -> Gold:
    book = load_workbook(path, read_only=True)
    rows = list(book["Hierarchy"].iter_rows(values_only=True))
    header = list(rows[0])
    label_col, parent_col = header.index("Label"), header.index("Parent")
    action_col = header.index("Action")
    company = rows[1][0]
    concepts: dict[str, Concept] = {}
    for row in rows[1:]:
        path_cells = [c for c in row[:label_col] if c]
        label, parent = row[label_col], row[parent_col]
        if path_cells[-1] != label:
            raise ValueError(f"level columns end with {path_cells[-1]!r}, label is {label!r}")
        actions = _actions([row[action_col], *_split(row[action_col + 1])])
        aliases = tuple(sorted(_split(row[action_col + 2])))
        prior = concepts.get(label)
        parents = tuple(sorted({parent, *(prior.parents if prior else ())}))
        concepts[label] = Concept(label, parents, actions, aliases)
    relations: set[Relation] = set()
    for r in list(book["Relations"].iter_rows(values_only=True))[1:]:
        relations.add(
            Relation(r[0], r[4], _actions([r[1], *_split(r[2])]), tuple(sorted(_split(r[3]))))
        )
    return Gold(company, concepts, relations)


def _read_json(path: Path) -> Gold:
    data = json.loads(path.read_text(encoding="utf-8"))
    company = data["label"]
    concepts: dict[str, Concept] = {}

    def walk(nodes: list[dict], parent: str) -> None:
        for n in nodes:
            parents = tuple(sorted([parent, *n.get("alsoUnder", [])]))
            concepts[n["label"]] = Concept(
                n["label"], parents, _actions(n["action"]), tuple(sorted(n.get("aliases", [])))
            )
            walk(n.get("children", []), n["label"])

    walk(data["children"], company)
    relations = {
        Relation(r["from"], r["to"], _actions(r["actions"]), tuple(sorted(r["inverse"])))
        for r in data["relations"]
    }
    return Gold(company, concepts, relations)


def _graph(ns: Namespace) -> Graph:
    g = Graph()
    g.bind("gold", ns)
    g.bind("oxe", OXE)
    g.bind("owl", OWL)
    g.bind("skos", SKOS)
    return g


def _iris(gold: Gold, ns: Namespace) -> dict[str, URIRef]:
    iri: dict[str, URIRef] = {gold.company: ns[_local(gold.company)]}
    for label in gold.concepts:
        name = _local(label)
        while ns[name] in iri.values():
            name += "_"
        iri[label] = ns[name]
    return iri


def _properties(gold: Gold, ns: Namespace, g: Graph) -> dict[Relation, URIRef]:
    ids = _verb_ids(gold)
    out = {}
    for r, tid in ids.items():
        prop = ns[tid]
        out[r] = prop
        if (prop, RDF.type, OWL.ObjectProperty) in g:
            continue
        g.add((prop, RDF.type, OWL.ObjectProperty))
        g.add((prop, RDFS.label, Literal(r.actions[0], lang="en")))
        for a in r.actions[1:]:
            g.add((prop, OXE.acceptedAction, Literal(a)))
        for a in r.inverse:
            g.add((prop, OXE.inverseAction, Literal(a)))
    return out


def _verb_ids(gold: Gold) -> dict[Relation, str]:
    """A property id per distinct (actions, inverse) pair, named after the first action."""
    by_signature: dict[tuple, str] = {}
    out = {}
    for r in sorted(gold.relations, key=str):
        signature = (r.actions, r.inverse)
        if signature not in by_signature:
            base = _local(r.actions[0], lower=True)
            name, n = base, 1
            while name in by_signature.values():
                n += 1
                name = f"{base}_{n}"
            by_signature[signature] = name
        out[r] = by_signature[signature]
    return out


def _add_actions(g: Graph, node: URIRef, actions: tuple[str, ...]) -> None:
    g.add((node, OXE.birthAction, Literal(actions[0])))
    for a in actions[1:]:
        g.add((node, OXE.acceptedAction, Literal(a)))


def _concept_of(g: Graph, node, label: str, parents: list[str]) -> Concept:
    birth = [str(o) for o in g.objects(node, OXE.birthAction)]
    accepted = [str(o) for o in g.objects(node, OXE.acceptedAction)]
    aliases = sorted(str(o) for o in g.objects(node, SKOS.altLabel))
    return Concept(label, tuple(sorted(parents)), _actions(birth + accepted), tuple(aliases))


def _relation_of(g: Graph, prop, source: str, target: str) -> Relation:
    first = str(g.value(prop, RDFS.label))
    accepted = [str(o) for o in g.objects(prop, OXE.acceptedAction)]
    inverse = sorted(str(o) for o in g.objects(prop, OXE.inverseAction))
    return Relation(source, target, _actions([first, *accepted]), tuple(inverse))


def _path(gold: Gold, label: str) -> list[str]:
    path = [label]
    while path[-1] != gold.company:
        path.append(gold.concepts[path[-1]].parents[0])
    return list(reversed(path))


def _actions(actions: list[str]) -> tuple[str, ...]:
    return (actions[0], *sorted(actions[1:])) if actions else ()


def _listed(value: object) -> list[str]:
    return [value] if isinstance(value, str) else list(value)


def _split(cell: object) -> list[str]:
    return [p.strip() for p in str(cell or "").split(";") if p.strip()]


def _local(label: str, lower: bool = False) -> str:
    parts = re.findall(r"[A-Za-z0-9]+", label)
    name = "".join(p[:1].upper() + p[1:] for p in parts)
    if lower:
        name = name[:1].lower() + name[1:]
    return name if name[:1].isalpha() else f"N{name}"


def _obo_quote(text: str) -> str:
    return text.replace("\\", "\\\\").replace('"', '\\"')


def _obo_values(lines: list[str]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for line in lines:
        m = re.match(r'(\S+) "((?:[^"\\]|\\.)*)" xsd:string', line)
        if m:
            out.setdefault(m.group(1), []).append(m.group(2).replace('\\"', '"'))
    return out


if __name__ == "__main__":
    sys.exit(main())
