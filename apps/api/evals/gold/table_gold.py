"""A hierarchy in a CSV file or an Excel workbook as a gold tree.

The header row is searched in the first 20 rows of each sheet (headers compare case-insensitively,
`_` as a space) and the first sheet with a recognised header is read. Three shapes:

- Edge list: a `parent` column and a `child` (or `label`, `name`, `concept`) column, optional
  `verb` (or `action`, `relation`), `accepted actions`, `definition` (or `description`) and
  `aliases` (or `synonyms`) columns; lists in a cell are separated by `;` or `|`. An empty
  parent, or the company's name, is the root. The verb defaults to "has".
- Hierarchical identifiers, such as the APQC Process Classification Framework: a
  `Hierarchy ID` column (1.0, 1.1, 1.1.1, ...) and a name column (`Name`, `Element`, `Process`).
  `N.0` (or `N`) is a top-level category under the root; the parent of 1.1.1 is 1.1, of 1.1 is
  1.0. The verb is the leading word of the name, lower-cased ("Develop Vision and Strategy" ->
  "develop"; "Develop and Manage Products" -> "develop" or "manage"), and the name without its
  verbs (and a leading article) is an alias. A row whose parent identifier is missing hangs off
  its nearest listed ancestor, else the root.
- Level columns: `Level 1`, `Level 2`, ... (and an optional `Company` column left of them, the
  root). A row's deepest filled level cell is a node; its parent is the nearest filled cell to
  its left, where a blank cell carries the value of the rows above. Optional verb, accepted
  actions, definition and aliases columns as for an edge list; the verb defaults to "has".

Relations come from another sheet of a workbook with `from` (or `source`), `to` (or `target`)
and `action` (or `verb`, `relation`) columns, optional `accepted actions` and `inverse actions`.
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from evals.gold.common import clean, finish, key, relation
from evals.gold.gold_tree import GoldTree
from evals.teach_case import ExpectedConcept, ExpectedRelation

TABLE_SUFFIXES = (".csv", ".tsv", ".xlsx", ".xlsm")
_EXCEL = (".xlsx", ".xlsm")
_HEADER_SEARCH_ROWS = 20

_PARENT = ("parent", "parent label", "parent name", "broader")
_CHILD = ("child", "label", "name", "concept", "child label", "child name", "term")
_VERB = ("verb", "action", "relation", "relationship", "birth action")
_ACCEPTED = ("accepted actions", "accepted action", "other actions")
_INVERSE = ("inverse actions", "inverse action", "inverse")
_DEFINITION = ("definition", "description", "element description", "process description")
_ALIASES = ("aliases", "alias", "synonyms", "synonym")
_HIERARCHY = ("hierarchy id", "hierarchy", "hierarchyid", "outline", "wbs")
_ELEMENT = ("name", "element", "process", "process element", "element name", "process name")
_PCF_ID = ("pcf id", "pcfid", "element id")
_ROOT_COLUMN = ("company", "root", "level 0")
_SOURCE = ("from", "source", "subject", "source label")
_TARGET = ("to", "target", "object", "target label")
_LEVEL = re.compile(r"^(level|lvl|l)\s*(\d+)$")
_HIERARCHY_VALUE = re.compile(r"^\d+(\.\d+)*$")
_LEADING_ARTICLE = re.compile(r"^(the|a|an)\s+", re.IGNORECASE)
_LIST_SEPARATOR = re.compile(r"[;|]")


@dataclass
class _Table:
    sheet: str
    shape: str  # "edges", "hierarchy" or "levels"
    header_row: int  # 1-based
    columns: dict[str, int]
    rows: list[list[str]] = field(default_factory=list)
    # Level columns, left to right (the company column first when there is one).
    levels: list[int] = field(default_factory=list)
    root_column: bool = False


@dataclass
class _Read:
    concepts: list[ExpectedConcept] = field(default_factory=list)
    skipped: Counter[str] = field(default_factory=Counter)
    # Hierarchy identifiers whose parent identifier is not listed.
    reattached: list[str] = field(default_factory=list)


def load_table(path: Path, company: str) -> GoldTree:
    sheets = _read_sheets(path)
    table = _find_table(path, sheets)
    fmt = f"{'xlsx' if path.suffix.lower() in _EXCEL else 'csv'}-{table.shape}"
    if table.shape == "edges":
        read = _edge_concepts(table, company)
    elif table.shape == "levels":
        read = _level_concepts(table, company)
    else:
        read = _hierarchy_concepts(table, company)
    relations, relation_sheet, relation_skipped = _relations(sheets, table.sheet)
    report: dict[str, object] = {
        "sheet": table.sheet,
        "sheetsInFile": list(sheets),
        "headerRow": table.header_row,
        "columns": sorted(table.columns),
        "rowsRead": len(table.rows),
        "rowsSkipped": dict(sorted(read.skipped.items())),
        "rowsAttachedToAnAncestor": read.reattached,
        "hasDefinitions": "definition" in table.columns,
        "relationSheet": relation_sheet,
        "relationRowsSkipped": relation_skipped,
    }
    return finish(read.concepts, relations, [], report, company, fmt)


def pcf_definitions_text(path: Path) -> str | None:
    """Every row's "Name. Definition" as paragraphs, when the table has a definitions column
    and a name column; None otherwise."""
    table = _find_table(path, _read_sheets(path))
    if "definition" not in table.columns or "name" not in table.columns:
        return None
    paragraphs: list[str] = []
    for row in table.rows:
        name = clean(_cell(row, table.columns["name"]))
        definition = clean(_cell(row, table.columns["definition"]))
        if name and definition:
            paragraphs.append(f"{name.rstrip('.')}. {definition}")
        elif name:
            paragraphs.append(f"{name.rstrip('.')}.")
    return "\n\n".join(paragraphs) if paragraphs else None


def leading_verbs(name: str) -> tuple[list[str], str]:
    """The leading verb of a process name, lower-cased (both verbs of "Develop and Manage ..."),
    and the rest of the name without a leading article."""
    words = name.split()
    if len(words) < 2:
        return [w.lower() for w in words], ""
    verbs, start = [words[0].lower()], 1
    if len(words) > 3 and words[1].lower() == "and":
        verbs, start = [words[0].lower(), words[2].lower()], 3
    rest = _LEADING_ARTICLE.sub("", " ".join(words[start:])).strip()
    return verbs, rest


def _edge_concepts(table: _Table, company: str) -> _Read:
    read = _Read()
    for row in table.rows:
        child = clean(_cell(row, table.columns["name"]))
        if not child:
            read.skipped["no label"] += 1
            continue
        parent = clean(_cell(row, table.columns["parent"])) or company
        read.concepts.append(_concept(table, row, child, parent))
    return read


def _level_concepts(table: _Table, company: str) -> _Read:
    read = _Read()
    carried: list[str] = [""] * len(table.levels)
    for row in table.rows:
        cells = [clean(_cell(row, col)) for col in table.levels]
        filled = [i for i, cell in enumerate(cells) if cell]
        if not filled:
            read.skipped["empty"] += 1
            continue
        deepest = filled[-1]
        for i in range(deepest + 1):
            if cells[i]:
                carried[i] = cells[i]
        carried[deepest + 1 :] = [""] * (len(carried) - deepest - 1)
        if table.root_column and deepest == 0:
            read.skipped["company row"] += 1
            continue
        above = [i for i in range(deepest) if carried[i]]
        if not above or (table.root_column and above[-1] == 0):
            parent = company
        else:
            parent = carried[above[-1]]
        read.concepts.append(_concept(table, row, carried[deepest], parent))
    return read


def _hierarchy_concepts(table: _Table, company: str) -> _Read:
    cols = table.columns
    read = _Read()
    names: dict[str, str] = {}
    order: list[str] = []
    for row in table.rows:
        raw_id, name = _cell(row, cols["hierarchy"]).strip(), clean(_cell(row, cols["name"]))
        if not raw_id and not name:
            read.skipped["empty"] += 1
            continue
        if not _HIERARCHY_VALUE.match(raw_id):
            read.skipped["no hierarchy id"] += 1
            continue
        if not name:
            read.skipped["no name"] += 1
            continue
        hid = _canonical(raw_id)
        if hid in names:
            read.skipped["duplicate hierarchy id"] += 1
            continue
        names[hid] = name
        order.append(hid)

    for hid in order:
        parts = hid.split(".")[:-1]
        if parts and ".".join(parts) not in names:
            read.reattached.append(hid)
        while parts and ".".join(parts) not in names:
            parts = parts[:-1]
        parent = names[".".join(parts)] if parts else company
        verbs, alias = leading_verbs(names[hid])
        aliases = [alias] if alias and key(alias) != key(names[hid]) else []
        read.concepts.append(
            ExpectedConcept(label=names[hid], parent=[parent], action=verbs, aliases=aliases)
        )
    return read


def _concept(table: _Table, row: list[str], label: str, parent: str) -> ExpectedConcept:
    """A concept of an edge-list or level-column row, with the row's verb columns and aliases."""
    cols = table.columns
    verb = clean(_cell(row, cols["verb"])).lower() if "verb" in cols else ""
    accepted = _listed(_cell(row, cols["accepted"])) if "accepted" in cols else []
    aliases = _listed(_cell(row, cols["aliases"])) if "aliases" in cols else []
    return ExpectedConcept(
        label=label,
        parent=[parent],
        action=[verb or "has", *(a.lower() for a in accepted)],
        aliases=aliases,
    )


def _relations(
    sheets: dict[str, list[list[str]]], tree_sheet: str
) -> tuple[list[ExpectedRelation], str | None, int]:
    """The relations of the first other sheet with a from/action/to header, its name, and the
    number of its non-empty rows skipped for a missing cell."""
    for sheet, rows in sheets.items():
        if sheet == tree_sheet:
            continue
        for index, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
            headers = _headers(row)
            cols = {
                role: col
                for role, names in (
                    ("source", _SOURCE),
                    ("target", _TARGET),
                    ("verb", _VERB),
                    ("accepted", _ACCEPTED),
                    ("inverse", _INVERSE),
                )
                if (col := _first(headers, names)) is not None
            }
            if not {"source", "target", "verb"} <= set(cols):
                continue
            relations, skipped = _relation_rows(rows[index + 1 :], cols)
            return relations, sheet, skipped
    return [], None, 0


def _relation_rows(
    rows: list[list[str]], cols: dict[str, int]
) -> tuple[list[ExpectedRelation], int]:
    relations: list[ExpectedRelation] = []
    skipped = 0
    for row in rows:
        source = clean(_cell(row, cols["source"]))
        target = clean(_cell(row, cols["target"]))
        verb = clean(_cell(row, cols["verb"])).lower()
        if not (source and target and verb):
            skipped += any(clean(c) for c in row)
            continue
        accepted = _listed(_cell(row, cols["accepted"])) if "accepted" in cols else []
        inverse = _listed(_cell(row, cols["inverse"])) if "inverse" in cols else []
        relations.append(
            relation(
                source,
                target,
                [verb, *(a.lower() for a in accepted)],
                [i.lower() for i in inverse],
            )
        )
    return relations, skipped


def _canonical(hid: str) -> str:
    """`1.0` and `1` are the same top-level identifier."""
    parts = hid.split(".")
    if len(parts) == 2 and parts[1] == "0":
        return parts[0]
    return hid


def _cell(row: list[str], index: int) -> str:
    return row[index] if index < len(row) else ""


def _listed(cell: str) -> list[str]:
    return [clean(part) for part in _LIST_SEPARATOR.split(cell) if clean(part)]


def _find_table(path: Path, sheets: dict[str, list[list[str]]]) -> _Table:
    for sheet, rows in sheets.items():
        table = _table_in(sheet, rows)
        if table is not None:
            return table
    raise ValueError(
        f"{path.name}: no sheet has a parent/child, hierarchy-id/name or level-column header "
        f"in its first {_HEADER_SEARCH_ROWS} rows"
    )


def _headers(row: list[str]) -> list[str]:
    return [" ".join(h.strip().lower().replace("_", " ").split()) for h in row]


def _table_in(sheet: str, rows: list[list[str]]) -> _Table | None:
    for index, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        headers = _headers(row)
        columns = {
            role: col
            for role, names in (
                ("parent", _PARENT),
                ("verb", _VERB),
                ("accepted", _ACCEPTED),
                ("definition", _DEFINITION),
                ("aliases", _ALIASES),
            )
            if (col := _first(headers, names)) is not None
        }
        body = rows[index + 1 :]
        hierarchy = _first(headers, _HIERARCHY)
        element = _first(headers, _ELEMENT)
        if hierarchy is not None and element is not None:
            columns = {k: v for k, v in columns.items() if k in ("definition", "aliases")}
            columns |= {"hierarchy": hierarchy, "name": element}
            pcf_id = _first(headers, _PCF_ID)
            if pcf_id is not None:
                columns["pcf id"] = pcf_id
            return _Table(sheet, "hierarchy", index + 1, columns, body)
        child = _first(headers, _CHILD)
        if "parent" in columns and child is not None:
            return _Table(sheet, "edges", index + 1, columns | {"name": child}, body)
        levels = sorted(
            (int(m.group(2)), col) for col, h in enumerate(headers) if (m := _LEVEL.match(h))
        )
        if len(levels) >= 2:
            level_cols = [col for _, col in levels]
            root = _first(headers, _ROOT_COLUMN)
            has_root = root is not None and root < level_cols[0]
            if root is not None and has_root:
                level_cols.insert(0, root)
            columns.pop("parent", None)
            return _Table(
                sheet, "levels", index + 1, columns, body, levels=level_cols, root_column=has_root
            )
    return None


def _first(headers: list[str], names: tuple[str, ...]) -> int | None:
    for name in names:
        if name in headers:
            return headers.index(name)
    return None


def _read_sheets(path: Path) -> dict[str, list[list[str]]]:
    if path.suffix.lower() in _EXCEL:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return {
                ws.title: [[_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
                for ws in workbook.worksheets
            }
        finally:
            workbook.close()
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], ",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    return {path.stem: [list(r) for r in csv.reader(text.splitlines(), dialect)]}


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return f"{value:.1f}"
    return str(value).strip()
