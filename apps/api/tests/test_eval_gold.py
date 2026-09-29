"""The gold-tree importers of the teach bake-off: every format read into expected concepts,
relations and optional leaves, and the format detected from the suffix and the content."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook

from evals.case_sources import benchmark_folders, discover
from evals.gold.gold_tree import GoldTree
from evals.gold.loader import GOLD_SUFFIXES, detect_format, load_gold, main
from evals.gold.table_gold import leading_verbs, pcf_definitions_text

FIXTURES = Path(__file__).resolve().parents[1] / "evals" / "fixtures" / "gold"
COMPANY = "Acme Corp"
OWL_SAMPLES = ("sample.ttl", "sample.rdf", "sample.owx", "sample.jsonld", "sample.nt")


def _concepts(tree: GoldTree) -> dict[str, tuple[list[str], list[str], list[str]]]:
    return {c.label: (c.parent, c.action, c.aliases) for c in tree.expected.concepts}


def _relations(tree: GoldTree) -> set[tuple[str, str, str]]:
    return {(r.source, r.action[0], r.target) for r in tree.expected.relations}


@pytest.mark.parametrize("name", OWL_SAMPLES)
def test_owl_samples_give_the_same_tree(name: str) -> None:
    tree = load_gold(FIXTURES / name, COMPANY)
    concepts = _concepts(tree)

    assert set(concepts) == {
        "Sales",
        "Operations",
        "Customer",
        "Order",
        "Online Order",
        "Express Online Order",
        "Delivery",
        "Product",
    }
    # The class labelled with the company's name is the root; owl:Thing hangs off it too.
    assert concepts["Sales"][0] == [COMPANY]
    assert concepts["Product"][0] == [COMPANY]
    # Two superclasses are both acceptable parents.
    assert sorted(concepts["Delivery"][0]) == ["Operations", "Sales"]
    # No label: the local name, humanised.
    assert concepts["Express Online Order"][0] == ["Online Order"]
    assert all(action == ["is a"] for _, action, _ in concepts.values())
    assert tree.report["maxDepth"] == 4


@pytest.mark.parametrize("name", OWL_SAMPLES)
def test_owl_samples_give_relations_verbs_and_optional_individuals(name: str) -> None:
    tree = load_gold(FIXTURES / name, COMPANY)

    assert _relations(tree) == {
        # rdfs:label with its cardinality annotation dropped.
        ("Customer", "places", "Order"),
        # No label: the camelCase local name split into words.
        ("Delivery", "ships product", "Product"),
        # A union domain gives one relation per member.
        ("Order", "contains", "Product"),
        ("Delivery", "contains", "Product"),
        # A someValuesFrom restriction.
        ("Customer", "prefers product", "Product"),
        # An assertion between individuals.
        ("BigCo", "prefers product", "Widget"),
    }
    assert tree.optional == ["BigCo", "Widget"]


@pytest.mark.parametrize("name", OWL_SAMPLES)
def test_owl_samples_report_what_they_ignore(name: str) -> None:
    report = load_gold(FIXTURES / name, COMPANY).report

    ignored = report["ignoredAxioms"]
    assert isinstance(ignored, dict)
    disjoint = "DisjointClasses" if name.endswith(".owx") else "owl:disjointWith"
    assert ignored[disjoint] == 1
    assert report["datatypeProperties"] == ["Order Date"]
    assert report["datatypeRestrictions"] == ["Order . Order Date"]
    assert report["individuals"] == 2
    assert report["source"] == name


def test_skos_builds_the_tree_from_broader_and_narrower() -> None:
    tree = load_gold(FIXTURES / "sample_skos.ttl", COMPANY)
    concepts = _concepts(tree)

    assert tree.report["format"] == "skos/turtle"
    # Top concepts (hasTopConcept, topConceptOf, or no broader) hang off the root.
    assert concepts["Products"][0] == [COMPANY]
    assert concepts["Services"][0] == [COMPANY]
    # narrower on the parent and broader on the child both give the tree.
    assert concepts["Hardware"][0] == ["Products"]
    assert concepts["Laptops"][0] == ["Hardware"]
    assert sorted(concepts["Device Support"][0]) == ["Hardware", "Services"]
    # prefLabel in English; altLabel as aliases; the verb is "has".
    assert concepts["Products"][2] == ["Goods"]
    assert all(action == ["has"] for _, action, _ in concepts.values())
    assert _relations(tree) == {("Device Support", "is related to", "Laptops")}
    assert tree.report["maxDepth"] == 4
    assert tree.report["ignoredAxioms"] == {"dct:modified": 1, "skos:definition": 1}


def test_obo_reads_is_a_relationships_synonyms_and_skips_obsolete_terms() -> None:
    tree = load_gold(FIXTURES / "sample.obo", COMPANY)
    concepts = _concepts(tree)

    assert tree.report["format"] == "obo"
    assert "Fax Order" not in concepts
    assert tree.report["obsoleteTerms"] == ["Fax Order"]
    assert concepts["Sales"][0] == [COMPANY]
    assert concepts["Product"][0] == [COMPANY]
    assert sorted(concepts["Delivery"][0]) == ["Operations", "Sales"]
    assert concepts["Customer"][2] == ["client", "buyer"]
    assert _relations(tree) == {
        ("Customer", "places", "Order"),
        # No typedef: the relation id split into words.
        ("Delivery", "ships product", "Product"),
        # The typedef's name.
        ("Express Delivery", "part of", "Operations"),
    }
    assert tree.report["ignoredAxioms"] == {"def": 1}
    assert tree.report["maxDepth"] == 3


def test_edge_list_csv() -> None:
    tree = load_gold(FIXTURES / "sample_edges.csv", COMPANY)
    concepts = _concepts(tree)

    assert tree.report["format"] == "csv-edges"
    assert concepts["Sales"] == ([COMPANY], ["has"], [])
    assert concepts["Operations"] == ([COMPANY], ["runs"], ["Ops"])
    assert concepts["Customer"] == (["Sales"], ["serves"], ["Client", "Buyer"])
    assert concepts["Online Order"][1] == ["has"]
    assert concepts["Delivery"][0] == ["Operations", "Sales"]
    assert tree.report["rowsSkipped"] == {"no label": 1}
    assert tree.report["maxDepth"] == 3


def test_hierarchical_id_csv() -> None:
    tree = load_gold(FIXTURES / "sample_pcf.csv", COMPANY)
    concepts = _concepts(tree)

    assert tree.report["format"] == "csv-hierarchy"
    assert concepts["Develop Vision and Strategy"] == (
        [COMPANY],
        ["develop"],
        ["Vision and Strategy"],
    )
    assert concepts["Define the business concept and long-term vision"][0] == [
        "Develop Vision and Strategy"
    ]
    assert concepts["Assess the external environment"][2] == ["external environment"]
    assert concepts["Analyze and evaluate competition"][:2] == (
        ["Assess the external environment"],
        ["analyze", "evaluate"],
    )
    assert concepts["Orphan Process"][0] == [COMPANY]
    assert tree.report["rowsAttachedToAnAncestor"] == ["9.9.9"]
    assert tree.report["rowsSkipped"] == {"empty": 1}
    assert tree.report["maxDepth"] == 4


def test_pcf_definitions_text() -> None:
    text = pcf_definitions_text(FIXTURES / "sample_pcf.csv")

    assert text is not None
    paragraphs = text.split("\n\n")
    assert paragraphs[0] == (
        "Develop Vision and Strategy. Establish a direction and vision for the organization."
    )
    assert "Identify economic trends." in paragraphs


def test_pcf_definitions_text_without_a_definitions_column(tmp_path: Path) -> None:
    path = tmp_path / "plain.csv"
    path.write_text("Hierarchy ID,Name\n1.0,Develop Vision and Strategy\n", encoding="utf-8")

    assert pcf_definitions_text(path) is None


def test_leading_verbs() -> None:
    assert leading_verbs("Develop Vision and Strategy") == (["develop"], "Vision and Strategy")
    assert leading_verbs("Manage the budget") == (["manage"], "budget")
    assert leading_verbs("Develop and Manage Products") == (["develop", "manage"], "Products")
    assert leading_verbs("Plan") == (["plan"], "")


def _pcf_workbook(path: Path) -> None:
    workbook = Workbook()
    intro = workbook.active
    intro.title = "Introduction"
    intro.append(["About this framework"])
    intro.append(["Version", "7.4"])
    pcf = workbook.create_sheet("Combined")
    pcf.append(["Process Classification Framework, cross-industry"])
    pcf.append([])
    pcf.append(["PCF ID", "Hierarchy ID", "Name", "Element Description"])
    pcf.append([10002, "1.0", "Develop Vision and Strategy", "Set the direction."])
    pcf.append([10014, "1.1", "Define the business concept", "Assess the environment."])
    pcf.append([10017, "1.1.1", "Assess the external environment", None])
    pcf.append([10021, "1.1.1.1", "Analyze and evaluate competition", "Study competitors."])
    pcf.append([10003, "2.0", "Develop and Manage Products and Services", None])
    workbook.save(path)


def test_hierarchical_id_workbook(tmp_path: Path) -> None:
    path = tmp_path / "pcf.xlsx"
    _pcf_workbook(path)

    tree = load_gold(path, COMPANY)
    concepts = _concepts(tree)

    assert tree.report["format"] == "xlsx-hierarchy"
    assert tree.report["sheet"] == "Combined"
    assert tree.report["sheetsInFile"] == ["Introduction", "Combined"]
    assert tree.report["headerRow"] == 3
    assert tree.report["rowsRead"] == 5
    assert concepts["Analyze and evaluate competition"][0] == ["Assess the external environment"]
    assert concepts["Develop and Manage Products and Services"][0] == [COMPANY]
    assert tree.report["maxDepth"] == 4
    text = pcf_definitions_text(path)
    assert text is not None
    assert text.startswith("Develop Vision and Strategy. Set the direction.")


def test_edge_list_workbook(tmp_path: Path) -> None:
    path = tmp_path / "edges.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["parent", "label", "action"])
    sheet.append(["", "Sales", "has"])
    sheet.append(["Sales", "Customer", "serves"])
    workbook.save(path)

    tree = load_gold(path, COMPANY)

    assert tree.report["format"] == "xlsx-edges"
    assert _concepts(tree)["Customer"] == (["Sales"], ["serves"], [])


def test_a_table_without_a_known_header_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "other.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")

    with pytest.raises(ValueError, match="no sheet has"):
        load_gold(path, COMPANY)


def test_nested_json() -> None:
    tree = load_gold(FIXTURES / "sample_nested.json", COMPANY)
    concepts = _concepts(tree)

    assert tree.report["format"] == "json-nested"
    # The top node named after the company is the root.
    assert COMPANY not in concepts
    assert concepts["Sales"] == ([COMPANY], ["has"], [])
    assert concepts["Customer"] == (["Sales"], ["serves"], ["Client"])
    assert concepts["Order"][1] == ["takes", "receives"]
    assert concepts["Express Online Order"][0] == ["Online Order"]
    assert tree.report["nodesSkipped"] == {"no label": 1}
    assert tree.report["maxDepth"] == 4


def test_graph_json() -> None:
    tree = load_gold(FIXTURES / "sample_graph.json", COMPANY)
    concepts = _concepts(tree)

    assert tree.report["format"] == "json-graph"
    assert concepts["Sales"][0] == [COMPANY]
    assert concepts["Customer"] == (["Sales"], ["is a"], [])
    # narrower reads parent -> child.
    assert concepts["Order"][0] == ["Sales"]
    assert concepts["Delivery"][0] == ["Operations", "Sales"]
    assert concepts["Product"][0] == [COMPANY]
    assert _relations(tree) == {
        ("Customer", "places", "Order"),
        ("Delivery", "ships product", "Product"),
    }
    assert tree.report["edgesSkipped"] == {"unknown node": 1}


@pytest.mark.parametrize(
    ("name", "fmt"),
    [
        ("sample.ttl", "turtle"),
        ("sample.rdf", "rdf-xml"),
        ("sample.owx", "owl-xml"),
        ("sample.jsonld", "json-ld"),
        ("sample.nt", "n-triples"),
        ("sample.obo", "obo"),
        ("sample_edges.csv", "csv"),
        ("sample_nested.json", "json"),
        ("sample_graph.json", "json"),
    ],
)
def test_format_by_suffix(name: str, fmt: str) -> None:
    assert detect_format(FIXTURES / name) == fmt


@pytest.mark.parametrize(
    ("source", "renamed", "fmt"),
    [
        ("sample.owx", "ontology.owl", "owl-xml"),
        ("sample.rdf", "ontology.owl", "rdf-xml"),
        ("sample.rdf", "ontology.xml", "rdf-xml"),
        ("sample.owx", "ontology.xml", "owl-xml"),
        ("sample.ttl", "ontology.owl", "turtle"),
        ("sample.jsonld", "ontology.json", "json-ld"),
        ("sample_nested.json", "tree.json", "json"),
    ],
)
def test_format_by_content(tmp_path: Path, source: str, renamed: str, fmt: str) -> None:
    path = tmp_path / renamed
    shutil.copy(FIXTURES / source, path)

    assert detect_format(path) == fmt
    tree = load_gold(path, COMPANY)
    assert tree.report["format"].startswith(fmt)
    assert tree.expected.concepts


def test_unsupported_suffix_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "ontology.docx"
    path.write_text("x", encoding="utf-8")

    with pytest.raises(ValueError, match="unsupported"):
        load_gold(path, COMPANY)
    assert ".xlsx" in GOLD_SUFFIXES and ".owx" in GOLD_SUFFIXES


def test_cli_prints_the_summary(capsys: pytest.CaptureFixture[str]) -> None:
    main([str(FIXTURES / "sample.ttl"), "--company", COMPANY])

    out = capsys.readouterr().out
    assert "format:    turtle" in out
    assert "concepts:  8" in out
    assert "max depth: 4" in out
    assert '"owl:disjointWith": 1' in out


def test_verb_annotations_on_owl_classes_and_properties() -> None:
    tree = load_gold(FIXTURES / "sample_annotated.ttl", COMPANY)
    concepts = _concepts(tree)

    # oxe:birthAction first, then the accepted actions; no annotation keeps "is a".
    assert concepts["Sales"][1] == ["runs", "has", "operates"]
    assert concepts["Invoice"][1] == ["produces"]
    assert concepts["Order"][1] == ["is a"]
    assert concepts["Order"][2] == ["Purchase order"]
    relation = tree.expected.relations[0]
    assert (relation.source, relation.target) == ("Invoice", "Order")
    assert relation.action == ["bills", "charges"]
    assert relation.inverse == ["is billed by"]
    assert "oxe:birthAction" not in tree.report["ignoredAxioms"]
    assert "skos:altLabel" not in tree.report["ignoredAxioms"]


def test_verb_annotations_in_owl_xml(tmp_path: Path) -> None:
    path = tmp_path / "annotated.owx"
    path.write_text(
        """<?xml version="1.0"?>
<Ontology xmlns="http://www.w3.org/2002/07/owl#" ontologyIRI="http://example.com/acme">
    <Prefix name="" IRI="http://example.com/acme#"/>
    <Prefix name="oxe" IRI="https://example.org/ontaix/evals/vocab#"/>
    <Prefix name="skos" IRI="http://www.w3.org/2004/02/skos/core#"/>
    <SubClassOf><Class abbreviatedIRI=":Sales"/><Class abbreviatedIRI=":Acme"/></SubClassOf>
    <AnnotationAssertion><AnnotationProperty abbreviatedIRI="rdfs:label"/>
        <AbbreviatedIRI>:Acme</AbbreviatedIRI><Literal>Acme Corp</Literal></AnnotationAssertion>
    <AnnotationAssertion><AnnotationProperty abbreviatedIRI="oxe:birthAction"/>
        <AbbreviatedIRI>:Sales</AbbreviatedIRI><Literal>runs</Literal></AnnotationAssertion>
    <AnnotationAssertion><AnnotationProperty abbreviatedIRI="oxe:acceptedAction"/>
        <AbbreviatedIRI>:Sales</AbbreviatedIRI><Literal>has</Literal></AnnotationAssertion>
    <AnnotationAssertion><AnnotationProperty abbreviatedIRI="skos:altLabel"/>
        <AbbreviatedIRI>:Sales</AbbreviatedIRI><Literal>Selling</Literal></AnnotationAssertion>
</Ontology>
""",
        encoding="utf-8",
    )

    assert _concepts(load_gold(path, COMPANY))["Sales"] == ([COMPANY], ["runs", "has"], ["Selling"])


def test_verb_annotations_and_property_relations_in_skos() -> None:
    tree = load_gold(FIXTURES / "sample_annotated_skos.ttl", COMPANY)
    concepts = _concepts(tree)

    assert concepts["Sales"][1] == ["runs", "has"]
    assert concepts["Invoice"][1] == ["produces"]
    assert concepts["Order"][1] == ["has"]
    relations = {(r.source, r.target): (r.action, r.inverse) for r in tree.expected.relations}
    # The property triple gives the relation; its skos:related twin adds nothing.
    assert relations == {
        ("Invoice", "Order"): (["bills", "charges"], ["is billed by"]),
        ("Customer", "Invoice"): (["is related to"], []),
    }


def test_verb_annotations_in_obo() -> None:
    tree = load_gold(FIXTURES / "sample_annotated.obo", COMPANY)
    concepts = _concepts(tree)

    assert concepts["Sales"][1] == ["runs", "has", "operates"]
    assert concepts["Invoice"][1] == ["produces"]
    assert concepts["Order"][1] == ["is a"]
    relation = tree.expected.relations[0]
    assert (relation.source, relation.action, relation.target) == (
        "Invoice",
        ["bills", "charges"],
        "Order",
    )
    assert relation.inverse == ["is billed by"]
    assert tree.report["ignoredAxioms"] == {"property_value owner": 1}


def test_nested_json_relations_and_further_parents() -> None:
    tree = load_gold(FIXTURES / "sample_nested_relations.json", COMPANY)
    concepts = _concepts(tree)

    assert concepts["Sales"][1] == ["runs", "has"]
    assert concepts["Invoice"][:2] == (["Sales", "Finance"], ["produces"])
    relations = {(r.source, r.target): (r.action, r.inverse) for r in tree.expected.relations}
    assert relations == {
        ("Invoice", "Order"): (["bills", "charges"], ["is billed by"]),
        ("Order", "Finance"): (["is paid to"], []),
    }
    assert tree.report["nodesSkipped"] == {"relation without from, to or action": 1}


def test_level_columns_csv() -> None:
    tree = load_gold(FIXTURES / "sample_levels.csv", COMPANY)
    concepts = _concepts(tree)

    assert tree.report["format"] == "csv-levels"
    assert concepts["Sales"][:2] == ([COMPANY], ["runs", "has", "operates"])
    # A blank left cell carries the value of the rows above.
    assert concepts["Order"][:2] == (["Sales"], ["has"])
    assert concepts["Online Order"][:2] == (["Order"], ["is a"])
    assert concepts["Invoice"][:2] == (["Sales"], ["produces"])
    assert concepts["Delivery"][0] == ["Operations"]
    assert concepts["Express Delivery"][0] == ["Delivery"]
    assert tree.report["rowsSkipped"] == {"empty": 1}
    assert tree.report["maxDepth"] == 3
    text = pcf_definitions_text(FIXTURES / "sample_levels.csv")
    assert text is None  # a level table has no single name column


def test_level_columns_workbook_with_a_relations_sheet(tmp_path: Path) -> None:
    path = tmp_path / "levels.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Outline"
    sheet.append(["Company", "Level 1", "Level 2", "Level 3", "Action"])
    sheet.append([COMPANY, None, None, None, None])
    sheet.append([COMPANY, "Sales", None, None, "runs"])
    sheet.append([None, None, "Order", None, None])
    sheet.append([None, None, None, "Order line", "contains"])
    sheet.append([COMPANY, "Finance", None, None, "runs"])
    relations = workbook.create_sheet("Relations")
    relations.append(["From", "Action", "Accepted actions", "Inverse actions", "To"])
    relations.append(["Order", "is paid to", "goes to", "receives", "Finance"])
    relations.append(["Order", None, None, None, None])
    workbook.save(path)

    tree = load_gold(path, COMPANY)
    concepts = _concepts(tree)

    assert tree.report["format"] == "xlsx-levels"
    assert concepts["Sales"][:2] == ([COMPANY], ["runs"])
    assert concepts["Order"][0] == ["Sales"]
    assert concepts["Order line"][:2] == (["Order"], ["contains"])
    assert concepts["Finance"][0] == [COMPANY]
    assert tree.report["rowsSkipped"] == {"company row": 1}
    assert tree.report["relationSheet"] == "Relations"
    assert tree.report["relationRowsSkipped"] == 1
    relation = tree.expected.relations[0]
    assert (relation.source, relation.action, relation.inverse, relation.target) == (
        "Order",
        ["is paid to", "goes to"],
        ["receives"],
        "Finance",
    )


BENCHMARKS = Path(__file__).resolve().parents[1] / "evals" / "benchmarks"
# (folder, gold file, company, concepts, relations)
PUBLIC_BENCHMARKS = (
    ("prov-o", "prov-o.ttl", "Provenance Ontology", 30, 61),
    ("dcat", "dcat3.ttl", "Data Catalog Vocabulary", 11, 10),
    ("ssn", "ssn.ttl", "Sensor Network Ontology", 22, 41),
    ("time", "time.ttl", "Time Ontology", 20, 34),
    ("valueflows", "valueflows.jsonld", "ValueFlows", 40, 324),
)


@pytest.mark.parametrize(("folder", "name", "company", "concepts", "relations"), PUBLIC_BENCHMARKS)
def test_public_benchmark_gold_parses(
    folder: str, name: str, company: str, concepts: int, relations: int
) -> None:
    tree = load_gold(BENCHMARKS / folder / name, company)

    assert len(tree.expected.concepts) == concepts
    assert len(tree.expected.relations) == relations
    assert all(c.parent for c in tree.expected.concepts)


def test_public_benchmark_cases_are_discovered(tmp_path: Path) -> None:
    found = discover([], benchmark_folders(BENCHMARKS), tmp_path)
    cases = {c.id: c for c in found.cases}

    assert found.notes == []
    for folder, name, company, concepts, _ in PUBLIC_BENCHMARKS:
        document = Path(name).stem
        case = cases[f"benchmarks-{folder}-{document}.md"]
        assert case.company == company
        assert case.document == BENCHMARKS / folder / f"{document}.md"
        assert len(case.expected.concepts) == concepts
        assert case.document.stat().st_size <= 100 * 1024
