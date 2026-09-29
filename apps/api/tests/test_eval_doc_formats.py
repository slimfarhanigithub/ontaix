"""Bake-off document loaders: every format to plain text and an upload the import API accepts.

Fixtures are built in a temporary directory; the OCR client is exercised with a fake and with a
mocked HTTP transport, so nothing reaches the network.
"""

from __future__ import annotations

import base64
import io
import json
import zipfile
from pathlib import Path

import httpx
import openpyxl
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.utilities.document_text import MEDIA_TYPE_BY_EXTENSION, extract_sentences
from evals.doc_formats.loaded_document import LoadedDocument
from evals.doc_formats.loader import DOCUMENT_SUFFIXES, DocumentFormatError, load_document
from evals.doc_formats.ocr_client import (
    AzureMistralOcrClient,
    FakeOcrClient,
    OcrError,
    mistral_ocr_url,
)

pytestmark = pytest.mark.asyncio(loop_scope="session")

OCR_TEXT = "# Scanned Invoice\n\nThe supplier ships the order to the customer within ten days."


def docx(paragraphs: list[str]) -> bytes:
    body = "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in paragraphs)
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}</w:body></w:document>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", xml)
    return out.getvalue()


def pdf(streams: list[str]) -> bytes:
    """A minimal PDF with one content stream per page."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(streams)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(streams)} >>".encode())
    font = 3 + 2 * len(streams)
    for i, content in enumerate(streams):
        stream = content.encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % n + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def text_pdf(pages: list[str]) -> bytes:
    return pdf([f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET" for text in pages])


def scanned_pdf() -> bytes:
    """A page with drawing operators and no text, as a scan's text layer looks to pypdf."""
    return pdf(["0 0 0 RG 72 72 m 540 720 l S 100 100 400 600 re f"])


def pptx() -> bytes:
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Service Catalogue"
    body = first.placeholders[1].text_frame
    body.text = "Insight sells consulting services."
    body.add_paragraph().text = "Every service belongs to one practice."
    first.notes_slide.notes_text_frame.text = "Mention the three practices."
    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Practices"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for r, row in enumerate([["Practice", "Lead"], ["Data", "Ana"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def xlsx() -> bytes:
    workbook = openpyxl.Workbook()
    customers = workbook.active
    customers.title = "Customers"
    customers.append(["Name", "Country", None])
    customers.append([None, None, None])
    customers.append(["Contoso Limited", "France", 12])
    orders = workbook.create_sheet("Orders")
    orders.append(["Order", "Customer"])
    orders.append(["O-1", "Contoso Limited"])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def write(tmp_path: Path, name: str, data: bytes) -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    return path


def assert_importable(doc: LoadedDocument) -> None:
    media_type = MEDIA_TYPE_BY_EXTENSION[Path(doc.upload_name).suffix.lower()]
    sentences, extracted, _ = extract_sentences(doc.upload_bytes, media_type)
    assert sentences, doc.upload_name
    assert extracted > 0


async def test_document_suffixes() -> None:
    assert DOCUMENT_SUFFIXES == {
        ".txt", ".text", ".md", ".markdown", ".html", ".htm", ".docx", ".pdf", ".pptx", ".xlsx"
    }  # fmt: skip


async def test_txt_keeps_original_bytes_and_tolerates_bom(tmp_path: Path) -> None:
    data = "﻿A customer places many orders.\r\n\r\nAn order has one customer.\r\n".encode()
    doc = await load_document(write(tmp_path, "notes.txt", data))
    assert doc.format == "txt"
    assert doc.text == "A customer places many orders.\n\nAn order has one customer.\n"
    assert doc.upload_name == "notes.txt" and doc.upload_bytes == data
    assert doc.words == 10 and doc.ocr is None and doc.pages is None
    assert_importable(doc)


async def test_md_keeps_original_bytes(tmp_path: Path) -> None:
    data = b"# Orders\n\nAn order lists the products a customer buys.\n"
    doc = await load_document(write(tmp_path, "orders.markdown", data))
    assert doc.format == "md"
    assert doc.text.startswith("# Orders\n\nAn order lists")
    assert doc.upload_name == "orders.markdown" and doc.upload_bytes == data
    assert_importable(doc)


async def test_html_strips_navigation_and_uploads_markdown(tmp_path: Path) -> None:
    html = (
        "<!DOCTYPE html><html><head><title>T</title><script>var x = 1;</script></head><body>"
        "<nav><a href='/'>Home menu link</a></nav><h1>Products</h1>"
        "<p>A product is something the company sells.</p><script>track('page')</script>"
        "<ul><li>Every product has a price.</li></ul></body></html>"
    )
    doc = await load_document(write(tmp_path, "spec.htm", html.encode()))
    assert doc.format == "html"
    assert "Home menu link" not in doc.text and "track" not in doc.text
    assert doc.text.index("# Products") < doc.text.index("A product is something")
    assert "- Every product has a price." in doc.text
    assert doc.upload_name == "spec.md" and doc.upload_bytes == doc.text.encode("utf-8")
    assert_importable(doc)


async def test_docx_paragraphs_in_order_with_original_upload(tmp_path: Path) -> None:
    data = docx(["Customers", "A customer places orders.", "An order ships from a warehouse."])
    doc = await load_document(write(tmp_path, "model.docx", data))
    assert doc.format == "docx"
    assert doc.text == (
        "Customers\n\nA customer places orders.\n\nAn order ships from a warehouse.\n"
    )
    assert doc.upload_name == "model.docx" and doc.upload_bytes == data
    assert_importable(doc)


async def test_text_pdf_pages_in_order_with_original_upload(tmp_path: Path) -> None:
    data = text_pdf(
        ["A warehouse stores products for delivery.", "A carrier delivers every shipment."]
    )
    doc = await load_document(write(tmp_path, "ops.pdf", data), FakeOcrClient(OCR_TEXT))
    assert doc.format == "pdf" and doc.pages == 2 and doc.ocr is None
    assert doc.text.index("A warehouse stores") < doc.text.index("A carrier delivers")
    assert "\n\n" in doc.text
    assert doc.upload_name == "ops.pdf" and doc.upload_bytes == data
    assert_importable(doc)


async def test_scanned_pdf_goes_through_ocr(tmp_path: Path) -> None:
    data = scanned_pdf()
    ocr = FakeOcrClient(OCR_TEXT, deployment="mistral-ocr-4-0")
    doc = await load_document(write(tmp_path, "invoice.pdf", data), ocr)
    assert ocr.calls == [data]
    assert doc.format == "scanned_pdf"
    assert doc.text == OCR_TEXT + "\n"
    assert doc.pages == 1
    assert doc.ocr is not None
    assert doc.ocr.deployment == "mistral-ocr-4-0" and doc.ocr.pages == 1
    assert doc.upload_name == "invoice.md" and doc.upload_bytes == doc.text.encode("utf-8")
    assert_importable(doc)


async def test_scanned_pdf_without_ocr_fails_clearly(tmp_path: Path) -> None:
    with pytest.raises(DocumentFormatError, match="scanned.*OCR client"):
        await load_document(write(tmp_path, "invoice.pdf", scanned_pdf()))


async def test_pptx_slides_titles_bodies_tables_and_notes(tmp_path: Path) -> None:
    doc = await load_document(write(tmp_path, "deck.pptx", pptx()))
    assert doc.format == "pptx" and doc.pages == 2
    expected_order = [
        "## Slide 1: Service Catalogue",
        "Insight sells consulting services.",
        "Every service belongs to one practice.",
        "Notes: Mention the three practices.",
        "## Slide 2: Practices",
        "Practice | Lead\nData | Ana",
    ]
    positions = [doc.text.index(fragment) for fragment in expected_order]
    assert positions == sorted(positions)
    assert doc.text.count("Service Catalogue") == 1
    assert doc.upload_name == "deck.md" and doc.upload_bytes == doc.text.encode("utf-8")
    assert_importable(doc)


async def test_xlsx_sheets_and_rows(tmp_path: Path) -> None:
    doc = await load_document(write(tmp_path, "book.xlsx", xlsx()))
    assert doc.format == "xlsx"
    assert doc.text == (
        "## Customers\n\nName | Country\nContoso Limited | France | 12\n\n"
        "## Orders\n\nOrder | Customer\nO-1 | Contoso Limited\n"
    )
    assert doc.upload_name == "book.md" and doc.upload_bytes == doc.text.encode("utf-8")
    assert_importable(doc)


@pytest.mark.parametrize(
    ("name", "data", "reason"),
    [
        ("fake.docx", b"just some text", "not a docx"),
        ("fake.pptx", docx(["A Word file named as a deck."]), "not a pptx"),
        ("fake.pdf", b"plain text pretending to be a PDF", "not a PDF"),
        ("fake.txt", b"%PDF-1.4 binary", "not txt"),
        ("fake.html", b"no markup in this file at all", "no HTML tags"),
        ("image.png", b"\x89PNG", "unsupported"),
    ],
)
async def test_content_must_match_extension(
    tmp_path: Path, name: str, data: bytes, reason: str
) -> None:
    with pytest.raises(DocumentFormatError, match=reason):
        await load_document(write(tmp_path, name, data))


async def test_mistral_ocr_url_from_foundry_endpoint() -> None:
    expected = "https://ontaix-dev.services.ai.azure.com/providers/mistral/azure/ocr"
    assert mistral_ocr_url("https://ontaix-dev.cognitiveservices.azure.com/") == expected
    assert mistral_ocr_url("https://ontaix-dev.services.ai.azure.com") == expected
    with pytest.raises(ValueError):
        mistral_ocr_url("https://example.com")


async def test_azure_mistral_ocr_request_and_response() -> None:
    seen: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request)
        return httpx.Response(
            200,
            json={
                "pages": [
                    {"index": 0, "markdown": "# Page One\n\nFirst page text."},
                    {"index": 1, "markdown": "Second page text."},
                ],
                "model": "mistral-ocr-4-0",
                "usage_info": {"pages_processed": 2, "doc_size_bytes": 10},
            },
        )

    tokens: list[str] = []

    def token_provider() -> str:
        tokens.append("issued")
        return "test-token"

    client = AzureMistralOcrClient.from_foundry_endpoint(
        "https://ontaix-dev.cognitiveservices.azure.com",
        deployment="mistral-ocr-4-0",
        price_eur_per_1000_pages=3.0,
        timeout_s=30,
        token_provider=token_provider,
        transport=httpx.MockTransport(handler),
    )
    result = await client.ocr_pdf(b"%PDF-1.4 scanned bytes")

    assert tokens == ["issued"]
    [request] = seen
    assert request.method == "POST"
    assert str(request.url) == (
        "https://ontaix-dev.services.ai.azure.com/providers/mistral/azure/ocr"
    )
    assert request.headers["Authorization"] == "Bearer test-token"
    body = json.loads(request.content)
    assert body == {
        "model": "mistral-ocr-4-0",
        "document": {
            "type": "document_url",
            "document_url": "data:application/pdf;base64,"
            + base64.b64encode(b"%PDF-1.4 scanned bytes").decode(),
        },
        "include_image_base64": False,
    }
    assert result.text == "# Page One\n\nFirst page text.\n\nSecond page text."
    assert result.pages == 2
    assert result.cost_eur == pytest.approx(0.006)
    assert result.latency_ms >= 0


async def test_azure_mistral_ocr_counts_pages_without_usage_info() -> None:
    transport = httpx.MockTransport(
        lambda request: httpx.Response(200, json={"pages": [{"markdown": "Only page."}]})
    )
    client = AzureMistralOcrClient(
        url="https://ocr.example.test/ocr",
        deployment="mistral-document-ai-2512",
        price_eur_per_1000_pages=10.0,
        token_provider=lambda: "t",
        transport=transport,
    )
    result = await client.ocr_pdf(b"%PDF-")
    assert result.text == "Only page." and result.pages == 1
    assert result.cost_eur == pytest.approx(0.01)


async def test_azure_mistral_ocr_http_error() -> None:
    transport = httpx.MockTransport(lambda request: httpx.Response(408, text="timeout"))
    client = AzureMistralOcrClient(
        url="https://ocr.example.test/ocr",
        deployment="mistral-ocr-4-0",
        price_eur_per_1000_pages=1.0,
        token_provider=lambda: "t",
        transport=transport,
    )
    with pytest.raises(OcrError, match="HTTP 408"):
        await client.ocr_pdf(b"%PDF-")
