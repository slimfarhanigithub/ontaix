"""Export (ADR 0016): formats, content types, scope and roles, tenant isolation, limits,
streaming, the Word document, and the OWL round trip through ontology import."""

from __future__ import annotations

import io
import re
import uuid
import zipfile
from typing import Any

import httpx
import owlrl
import pytest
from docx import Document
from owlrl.Namespaces import ERRNS
from rdflib import BNode, Graph, URIRef
from rdflib.namespace import OWL, RDF, RDFS, SKOS

from app.clients import db_client
from app.config import get_settings
from app.repositories import tenant_repository, tenant_settings_repository, view_state_repository
from app.routers import export as export_router
from app.seed import northwind
from app.services import company_service
from app.services.ontology_view_service import load_view
from app.utilities.ontaix_vocabulary import OX
from tests.conftest import TenantFixture
from tests.test_proposals import add_company

pytestmark = pytest.mark.asyncio(loop_scope="session")

RDF_FORMATS = {"owl": "xml", "turtle": "turtle", "jsonld": "json-ld", "skos": "turtle"}
MEDIA_TYPES = {
    "owl": "application/rdf+xml",
    "owx": "application/owl+xml",
    "turtle": "text/turtle",
    "jsonld": "application/ld+json",
    "skos": "text/turtle",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
EXTENSIONS = {
    "owl": ".owl",
    "owx": ".owx",
    "turtle": ".ttl",
    "jsonld": ".jsonld",
    "skos": ".skos.ttl",
    "docx": ".docx",
}
UUID_TEXT = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}")
TAUGHT = (
    ("Machine", "rated power", "number", "75.5"),
    ("Machine", "commissioned", "date", "2021-04-01"),
    ("Plant", "colour code", "text", "RAL 5015"),
)


async def export(
    client: httpx.AsyncClient, persona: Any, fmt: str, scope: str = "company", **ids: str
) -> httpx.Response:
    params = {"scope": scope, "format": fmt, **ids}
    return await client.get("/export", params=params, headers=persona.headers)


async def approve_all(client: httpx.AsyncClient, tenant: TenantFixture) -> None:
    response = await client.post("/proposals/approve-all", headers=tenant.governor.headers)
    assert response.status_code == 200, response.text


async def propose(client: httpx.AsyncClient, tenant: TenantFixture, drafts: list[dict]) -> None:
    for start in range(0, len(drafts), 200):
        response = await client.post(
            "/proposals/batch",
            json={"drafts": drafts[start : start + 200]},
            headers=tenant.builder.headers,
        )
        assert response.status_code == 202, response.text


async def build_model(client: httpx.AsyncClient, tenant: TenantFixture) -> dict:
    """The Northwind fixture model in the tenant's home company, with a reversed birth, a
    further `is a`, taught attributes, a pending concept and a second company linked by an
    equivalence; everything but the pending concept approved."""
    company = str(tenant.company_id)
    for batch in northwind.BATCHES:
        drafts = []
        for row in batch.rows:
            if isinstance(row, northwind.RelationRow):
                drafts.append(
                    {
                        "type": "relation",
                        "companyId": company,
                        "aLabel": row.subject,
                        "bLabel": row.object,
                        "action": row.action,
                    }
                )
                continue
            draft: dict[str, Any] = {
                "companyId": company,
                "label": row.label,
                "domainKey": row.domain,
            }
            if row.parent == northwind.COMPANY_NAME:
                draft["parentId"] = str(tenant.root_id)
            else:
                draft["parentLabel"] = row.parent
            if isinstance(row, northwind.SpecRow):
                draft.update({"type": "spec", "rule": row.rule})
            else:
                draft.update({"type": "concept", "action": row.action})
            drafts.append(draft)
        await propose(client, tenant, drafts)
        await approve_all(client, tenant)
    extras: list[dict] = [
        {
            "type": "concept",
            "companyId": company,
            "parentLabel": "Machine",
            "label": "Maintenance crew",
            "domainKey": "maintenance",
            "action": "services",
            "reverse": True,
        },
        {
            "type": "relation",
            "companyId": company,
            "aLabel": "Maintenance crew",
            "bLabel": "Operator",
            "action": "is a",
        },
    ]
    for label, name, attribute_type, value in TAUGHT:
        extras.append(
            {
                "type": "attr",
                "conceptLabel": label,
                "companyId": company,
                "name": name,
                "attributeType": attribute_type,
                "value": value,
            }
        )
    await propose(client, tenant, extras)
    await approve_all(client, tenant)
    other = await add_company(client, tenant, "Aurora")
    await propose(
        client,
        tenant,
        [
            {
                "type": "concept",
                "companyId": other["company"]["id"],
                "parentId": other["root"]["id"],
                "label": "Valve",
                "domainKey": "production",
                "action": "makes",
            }
        ],
    )
    await approve_all(client, tenant)
    scene = await read_scene(client, tenant)
    valve = next(n for n in scene["nodes"] if n["label"] == "Valve")
    machine = next(
        n for n in scene["nodes"] if n["label"] == "Machine" and n["companyId"] == company
    )
    equivalence = await client.post(
        "/equivalences",
        json={"aId": machine["id"], "bId": valve["id"]},
        headers=tenant.builder.headers,
    )
    assert equivalence.status_code == 202, equivalence.text
    await approve_all(client, tenant)
    await propose(
        client,
        tenant,
        [
            {
                "type": "concept",
                "companyId": company,
                "parentId": str(tenant.root_id),
                "label": "Pending thing",
                "domainKey": "sales",
                "action": "has",
            }
        ],
    )
    return other


async def read_scene(client: httpx.AsyncClient, tenant: TenantFixture) -> dict:
    response = await client.get("/scene", headers=tenant.governor.headers)
    assert response.status_code == 200, response.text
    return response.json()


def company_model(scene: dict, company_id: str) -> dict[str, Any]:
    """The approved model of one company, by label, with no ids: each concept's parent, birth,
    domain and rule; the relations that are not births within the company; taught attributes."""
    nodes = {n["id"]: n for n in scene["nodes"] if not n["pending"]}
    links = {link["id"]: link for link in scene["links"] if not link["pending"]}
    members = {i: n for i, n in nodes.items() if n["companyId"] == company_id}
    births = {n["birthRelationId"] for n in members.values() if n.get("birthRelationId")}
    concepts = {}
    for node in members.values():
        if node["kind"] == "root":
            continue
        birth = links[node["birthRelationId"]]
        parent = nodes[node["parentId"]]
        spec = birth["kind"] == "isa"
        concepts[node["label"]] = {
            "parent": "<root>" if parent["kind"] == "root" else parent["label"],
            "spec": spec,
            "action": None if spec else birth["label"],
            "reverse": (not spec) and birth["aId"] == node["id"],
            "domain": node["domainKey"],
            "rule": node["rule"] if spec else None,
        }
    relations = {
        (nodes[link["aId"]]["label"], link["label"], nodes[link["bId"]]["label"])
        for i, link in links.items()
        if i not in births
        and link["kind"] in ("rel", "isa")
        and link["aId"] in members
        and link["bId"] in members
    }
    attributes = {
        (node["label"], a["name"], a["type"], a["value"])
        for node in members.values()
        for a in node["attributes"]
        if a["value"] is not None and a["state"] == "approved"
    }
    return {"concepts": concepts, "relations": relations, "attributes": attributes}


async def test_every_format_has_its_content_type_file_name_and_parses_back(
    client: httpx.AsyncClient, tenant: TenantFixture
) -> None:
    await build_model(client, tenant)
    for fmt in MEDIA_TYPES:
        response = await export(client, tenant.builder, fmt, companyId=str(tenant.company_id))

        assert response.status_code == 200, response.text
        assert response.headers["content-type"].split(";")[0] == MEDIA_TYPES[fmt]
        assert response.headers["cache-control"] == "no-store"
        disposition = response.headers["content-disposition"]
        assert disposition.startswith('attachment; filename="ontaix-home-')
        assert disposition.endswith(f'{EXTENSIONS[fmt]}"')
        assert re.search(r"-\d{4}-\d{2}-\d{2}\.", disposition)
        if fmt in RDF_FORMATS:
            graph = Graph().parse(data=response.content, format=RDF_FORMATS[fmt])
            assert len(graph) > 100
            if fmt == "skos":
                labels = {str(o) for o in graph.objects(None, SKOS.prefLabel)}
            else:
                labels = {str(o) for o in graph.objects(None, RDFS.label)}
            assert {"Plant", "Machine due for maintenance"} <= labels
            assert "Pending thing" not in labels


async def test_owl_formats_carry_the_same_axioms_and_validate(
    client: httpx.AsyncClient, tenant: TenantFixture
) -> None:
    await build_model(client, tenant)
    company = str(tenant.company_id)
    owl = Graph().parse(
        data=(await export(client, tenant.builder, "owl", companyId=company)).content,
        format="xml",
    )
    turtle = Graph().parse(
        data=(await export(client, tenant.builder, "turtle", companyId=company)).content,
        format="turtle",
    )
    jsonld = Graph().parse(
        data=(await export(client, tenant.builder, "jsonld", companyId=company)).content,
        format="json-ld",
    )
    for graph in (turtle, jsonld):
        assert _ground(graph) == _ground(owl)
        assert len(graph) == len(owl)
    classes = {
        str(label): cls
        for cls in owl.subjects(RDF.type, OWL.Class)
        for label in owl.objects(cls, RDFS.label)
    }
    spec, parent = classes["Machine due for maintenance"], classes["Machine"]
    assert (spec, RDFS.subClassOf, parent) in owl
    line, plant = classes["Production line"], classes["Plant"]
    assert (line, RDFS.subClassOf, plant) not in owl
    assert (line, OX.bornFrom, plant) in owl
    assert str(owl.value(line, OX.birthAction)) == "runs"
    assert str(owl.value(line, OX.domain)) == "production"
    restrictions = {
        (str(owl.value(owl.value(r, OWL.onProperty), RDFS.label)), owl.value(r, OWL.someValuesFrom))
        for r in owl.objects(plant, RDFS.subClassOf)
        if (r, RDF.type, OWL.Restriction) in owl
    }
    assert ("runs", line) in restrictions
    same = list(owl.objects(classes["Machine"], OWL.equivalentClass))
    assert len(same) == 1 and "Valve" not in classes
    owlrl.DeductiveClosure(owlrl.OWLRL_Semantics).expand(owl)
    assert not list(owl.subjects(RDF.type, ERRNS.ErrorMessage))


async def test_skos_has_schemes_broader_collections_and_notes(
    client: httpx.AsyncClient, tenant: TenantFixture
) -> None:
    await build_model(client, tenant)
    response = await export(client, tenant.builder, "skos", scope="all")

    graph = Graph().parse(data=response.content, format="turtle")
    schemes = list(graph.subjects(RDF.type, SKOS.ConceptScheme))
    assert len(schemes) == 2
    concepts = {
        str(graph.value(c, SKOS.prefLabel)): c for c in graph.subjects(RDF.type, SKOS.Concept)
    }
    assert (concepts["Production line"], SKOS.broader, concepts["Plant"]) in graph
    assert (concepts["Plant"], SKOS.topConceptOf, None) in graph
    assert (concepts["Machine"], SKOS.exactMatch, concepts["Valve"]) in graph
    notes = {str(n) for n in graph.objects(concepts["Machine"], SKOS.note)}
    assert "rated power: 75.5" in notes
    assert list(graph.subjects(RDF.type, SKOS.Collection))


async def test_word_document_describes_the_model_in_prose_and_tables(
    client: httpx.AsyncClient, tenant: TenantFixture
) -> None:
    await build_model(client, tenant)
    response = await export(client, tenant.builder, "docx", scope="all")

    assert response.status_code == 200, response.text
    document = Document(io.BytesIO(response.content))
    headings = [
        (p.style.name, p.text)
        for p in document.paragraphs
        if p.style.name.startswith(("Heading", "Title"))
    ]
    assert ("Title", "Ontology Builder export") in headings
    assert ("Heading 1", tenant.company_name) in headings
    for name in ("Domains", "Entity hierarchy", "Relationships", "Equivalences"):
        assert ("Heading 2", name) in headings
    text = "\n".join(p.text for p in document.paragraphs)
    cells = "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "Machine due for maintenance" in text and "is a Machine" in text
    assert "rated power: 75.5" in text
    assert "Pending thing" not in text + cells
    assert "Production" in cells and "#d30c55" in cells
    assert not UUID_TEXT.search(text + cells)
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        names = archive.namelist()
        body = archive.read("word/document.xml").decode()
    assert not any("vba" in n.lower() for n in names)
    assert "w:fldChar" not in body and "w:instrText" not in body and "w:hyperlink" not in body
    assert 'w:fill="D30C55"' in body


@pytest.mark.parametrize("fmt", ["owl", "owx", "turtle", "jsonld"])
async def test_an_owl_export_imports_back_into_the_same_model(
    client: httpx.AsyncClient, tenant: TenantFixture, fmt: str
) -> None:
    await build_model(client, tenant)
    source = company_model(await read_scene(client, tenant), str(tenant.company_id))
    exported = await export(client, tenant.builder, fmt, companyId=str(tenant.company_id))
    assert exported.status_code == 200, exported.text
    target = await add_company(client, tenant, f"Copy {fmt}")
    target_id = target["company"]["id"]

    imported = await client.post(
        "/ontology-imports",
        files={"file": (f"export{EXTENSIONS[fmt]}", exported.content, MEDIA_TYPES[fmt])},
        data={"companyId": target_id},
        headers=tenant.builder.headers,
    )
    assert imported.status_code == 200, imported.text
    result = imported.json()
    reasons = {s["reason"] for s in result["skipped"]}
    assert reasons <= {"equivalence_not_imported", "unknown_parent"}, result["skipped"]
    proposed = await client.post(
        f"/ontology-imports/{result['ontologyImportId']}/proposals",
        json={"indexes": list(range(len(result["drafts"])))},
        headers=tenant.builder.headers,
    )
    assert proposed.status_code == 202, proposed.text
    await approve_all(client, tenant)

    copy = company_model(await read_scene(client, tenant), target_id)
    assert copy["concepts"] == source["concepts"]
    assert copy["relations"] == source["relations"]
    assert copy["attributes"] == source["attributes"]


async def test_a_company_owner_exports_only_what_it_reads(
    client: httpx.AsyncClient, tenant: TenantFixture
) -> None:
    other = await build_model(client, tenant)

    everything = await export(client, tenant.owner, "turtle", scope="all")
    assert everything.status_code == 200, everything.text
    graph = Graph().parse(data=everything.content, format="turtle")
    labels = {str(o) for o in graph.objects(None, RDFS.label)}
    assert "Plant" in labels and "Valve" not in labels
    assert not list(graph.subject_objects(OWL.equivalentClass))
    assert other["company"]["name"] not in labels

    refused = await export(client, tenant.owner, "owl", companyId=other["company"]["id"])
    assert refused.status_code == 404, refused.text
    outsider = await export(client, tenant.outsider, "owl", scope="all")
    assert outsider.status_code == 403, outsider.text


async def test_another_tenants_company_and_domain_are_not_found(
    client: httpx.AsyncClient, tenant: TenantFixture
) -> None:
    async with db_client.platform_session() as s:
        foreign = await tenant_repository.create(s, f"x-{uuid.uuid4().hex[:8]}", "Foreign")
        await tenant_settings_repository.create(s, foreign.id)
        await view_state_repository.create(s, foreign.id)
        view = await load_view(s, foreign.id)
        company = await company_service.add_company(s, view, "Foreign Co", "", is_home=True)
        product = next(p for p in view.domain_products.values() if p.company_id == company.id)
        await s.commit()

    by_company = await export(client, tenant.governor, "owl", companyId=str(company.id))
    by_domain = await export(
        client, tenant.governor, "owl", scope="domain", domainProductId=str(product.id)
    )
    everything = await export(client, tenant.governor, "turtle", scope="all")

    assert by_company.status_code == 404 and by_domain.status_code == 404
    assert "Foreign Co" not in everything.text


async def test_a_domain_scope_exports_that_domain_product_only(
    client: httpx.AsyncClient, tenant: TenantFixture
) -> None:
    await build_model(client, tenant)
    products = (await client.get("/domain-products", headers=tenant.governor.headers)).json()
    maintenance = next(
        p
        for p in products
        if p["key"] == "maintenance" and p["companyId"] == str(tenant.company_id)
    )

    response = await export(
        client, tenant.builder, "turtle", scope="domain", domainProductId=maintenance["id"]
    )

    assert response.status_code == 200, response.text
    assert "-maintenance-" in response.headers["content-disposition"]
    graph = Graph().parse(data=response.content, format="turtle")
    domains = {str(o) for o in graph.objects(None, OX.domain)}
    assert domains == {"maintenance"}
    labels = {str(o) for o in graph.objects(None, RDFS.label)}
    assert "Machine due for maintenance" in labels and "Plant" not in labels


async def test_bad_parameters_limit_budget_and_audit(
    client: httpx.AsyncClient, tenant: TenantFixture, monkeypatch: pytest.MonkeyPatch
) -> None:
    unknown_format = await export(client, tenant.builder, "pdf", scope="all")
    missing_id = await export(client, tenant.builder, "owl", scope="company")
    assert unknown_format.status_code == 400 and missing_id.status_code == 400

    await propose(
        client,
        tenant,
        [
            {
                "type": "concept",
                "companyId": str(tenant.company_id),
                "parentId": str(tenant.root_id),
                "label": label,
                "domainKey": "sales",
                "action": "has",
            }
            for label in ("Order", "Quote")
        ],
    )
    await approve_all(client, tenant)
    monkeypatch.setenv("ONTAIX_EXPORT_MAX_CONCEPTS", "1")
    get_settings.cache_clear()
    try:
        too_large = await export(client, tenant.governor, "owl", scope="all")
    finally:
        monkeypatch.delenv("ONTAIX_EXPORT_MAX_CONCEPTS")
        get_settings.cache_clear()
    assert too_large.status_code == 413, too_large.text
    assert too_large.json()["code"] == "payload_too_large"

    monkeypatch.setenv("ONTAIX_EXPORT_PER_HOUR", "1")
    get_settings.cache_clear()
    try:
        first = await export(client, tenant.admin, "turtle", scope="all")
        second = await export(client, tenant.admin, "turtle", scope="all")
    finally:
        monkeypatch.delenv("ONTAIX_EXPORT_PER_HOUR")
        get_settings.cache_clear()
    assert first.status_code == 200, first.text
    assert second.status_code == 429 and second.json()["code"] == "rate_limited"
    assert "Retry-After" in second.headers

    audit = (await client.get("/audit?filter[kind]=export", headers=tenant.admin.headers)).json()
    entries = [e for e in audit["items"] if e["kind"] == "export"]
    assert entries and entries[0]["what"] == "every readable company exported as OWL 2 Turtle"
    assert "Order" not in entries[0]["what"]


async def test_a_large_scope_is_streamed_in_chunks(
    client: httpx.AsyncClient, tenant: TenantFixture
) -> None:
    await propose(
        client,
        tenant,
        [
            {
                "type": "concept",
                "companyId": str(tenant.company_id),
                "parentId": str(tenant.root_id),
                "label": f"Part family {n:03d}",
                "domainKey": "production",
                "action": "has",
            }
            for n in range(400)
        ],
    )
    await approve_all(client, tenant)

    chunks: list[bytes] = []
    async with client.stream(
        "GET",
        "/export",
        params={"scope": "company", "companyId": str(tenant.company_id), "format": "turtle"},
        headers=tenant.builder.headers,
    ) as response:
        assert response.status_code == 200
        assert "content-length" not in response.headers
        async for chunk in response.aiter_raw():
            chunks.append(chunk)

    body = b"".join(chunks)
    assert len(body) > export_router.CHUNK_BYTES
    assert len(list(export_router._chunks(body))) > 1
    graph = Graph().parse(data=body, format="turtle")
    labels = {str(o) for o in graph.objects(None, RDFS.label)}
    assert {f"Part family {n:03d}" for n in range(400)} <= labels


def _ground(graph: Graph) -> set[tuple[str, str, str]]:
    """The triples without blank nodes or the export time, for comparing serialisations of one
    graph."""
    return {
        (str(s), str(p), str(o))
        for s, p, o in graph
        if isinstance(s, URIRef) and not isinstance(o, BNode) and p != OWL.versionInfo
    }
