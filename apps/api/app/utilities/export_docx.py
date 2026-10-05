"""The Word document of an export snapshot: the model described for people.

A title page with the scope, the date and the exporting tenant's name; then per company its
domains (name, owner, colour as a shaded cell), the entity hierarchy as nested lists following
the birth tree (each entry its label, its domain, its `is a` parents and its attributes), a table
of relationships (subject, action, object) and its equivalences to other readable companies. No
id appears. Every text is a plain run, so no label can become markup or a field; the document
has no macros, fields or external links. List entries and table rows are appended as
WordprocessingML elements with style ids resolved once, so a large model is written in linear
time.
"""

from __future__ import annotations

import io
import re
import uuid
from collections import defaultdict

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from lxml.etree import SubElement, _Element

from app.models.export.export_snapshot import (
    ExportCompany,
    ExportConcept,
    ExportRelation,
    ExportSnapshot,
)

HEADING_DOMAINS = "Domains"
HEADING_HIERARCHY = "Entity hierarchy"
HEADING_RELATIONSHIPS = "Relationships"
HEADING_EQUIVALENCES = "Equivalences"
NONE_TEXT = "None."
# A list level's indent, in twentieths of a point (0.63 cm).
LIST_INDENT_TWIPS = 357
_HEX = re.compile(r"^#[0-9a-fA-F]{6}$")
# Characters XML 1.0 cannot carry; labels never hold them, and a run never receives them.
_NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")

Run = tuple[str, bool]


def export_docx(snapshot: ExportSnapshot) -> bytes:
    return _Writer(snapshot).write()


class _Writer:
    def __init__(self, snapshot: ExportSnapshot) -> None:
        self.s = snapshot
        self.doc: DocxDocument = Document()
        self.sect_pr = self.doc.element.body.find(qn("w:sectPr"))
        self.list_style = self.doc.styles["List Bullet"].style_id
        self.concepts: dict[uuid.UUID, ExportConcept] = {
            c.id: c for c in [*snapshot.outside, *snapshot.concepts]
        }
        self.in_scope = {c.id for c in snapshot.concepts}
        self.company_names = {
            c.id: c.name for c in [*snapshot.outside_companies, *snapshot.companies]
        }
        self.domain_names = {d.key: d.name for d in snapshot.domains}
        self.isa: dict[uuid.UUID, list[uuid.UUID]] = defaultdict(list)
        for relation in snapshot.relations:
            if relation.kind == "isa":
                self.isa[relation.a_id].append(relation.b_id)

    def write(self) -> bytes:
        self._properties()
        self._title_page()
        for company in self.s.companies:
            self._company(company)
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()

    def _properties(self) -> None:
        props = self.doc.core_properties
        props.title = _text(f"Ontology Builder export · {self.s.scope_name}")
        props.author = "Ontology Builder"
        props.last_modified_by = "Ontology Builder"
        props.created = self.s.exported_at
        props.modified = self.s.exported_at
        props.comments = ""

    def _title_page(self) -> None:
        doc, s = self.doc, self.s
        doc.add_heading("Ontology Builder export", level=0)
        for name, value in (
            ("Scope", s.scope_name),
            ("Date", s.exported_at.date().isoformat()),
            ("Tenant", s.tenant_name),
        ):
            self._paragraph([(f"{name}: ", True), (value, False)])
        self._paragraph(
            [("The approved model of the scope: pending proposals are not included.", False)]
        )
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _company(self, company: ExportCompany) -> None:
        self.doc.add_heading(_text(company.name), level=1)
        if company.sub:
            self._paragraph([(company.sub, False)])
        self._domains(company)
        self._hierarchy(company)
        self._relationships(company)
        self._equivalences(company)

    def _domains(self, company: ExportCompany) -> None:
        self.doc.add_heading(HEADING_DOMAINS, level=2)
        domains = [d for d in self.s.domains if d.key in company.domain_keys]
        if not domains:
            self._paragraph([(NONE_TEXT, False)])
            return
        table = self._table(("Domain", "Owner", "Colour"))
        for domain in domains:
            cells = self._row(table, (domain.name, domain.owner or "—", domain.color))
            _shade(cells[2], domain.color)

    def _hierarchy(self, company: ExportCompany) -> None:
        self.doc.add_heading(HEADING_HIERARCHY, level=2)
        members = [c for c in self.s.concepts if c.company_id == company.id and c.birth != "root"]
        children: dict[uuid.UUID | None, list[ExportConcept]] = defaultdict(list)
        for concept in members:
            parent = concept.parent_id
            top = parent is None or parent not in self.in_scope or self._is_root(parent)
            children[None if top else parent].append(concept)
        for group in children.values():
            group.sort(key=lambda c: c.label.lower())
        if not members:
            self._paragraph([(NONE_TEXT, False)])
            return
        # Depth first, without recursion: the tree is as deep as the model.
        stack: list[tuple[ExportConcept, int]] = [(c, 0) for c in reversed(children[None])]
        while stack:
            concept, depth = stack.pop()
            self._entry(concept, depth)
            stack.extend((c, depth + 1) for c in reversed(children.get(concept.id, [])))

    def _entry(self, concept: ExportConcept, depth: int) -> None:
        details: list[str] = []
        if concept.domain_key is not None:
            details.append(self.domain_names.get(concept.domain_key, concept.domain_key))
        parents = [self.concepts[i] for i in self.isa.get(concept.id, []) if i in self.concepts]
        if concept.birth == "spec" and concept.parent_id in self.concepts:
            parents.insert(0, self.concepts[concept.parent_id])
        if parents:
            details.append("is a " + ", ".join(self._name(p, concept) for p in parents))
        elif concept.birth == "birth" and concept.parent_id in self.concepts:
            parent = self._name(self.concepts[concept.parent_id], concept)
            action = concept.birth_action or ""
            details.append(
                f"{concept.label} {action} {parent}"
                if concept.birth_reverse
                else f"{parent} {action} {concept.label}"
            )
        if concept.rule:
            details.append(f"rule: {concept.rule}")
        for attribute in concept.attributes:
            if attribute.taught:
                details.append(f"{attribute.name}: {attribute.value}")
            else:
                source = attribute.col or "a source"
                details.append(f"{attribute.name} ({attribute.type}, from {source})")
        runs: list[Run] = [(concept.label, True)]
        if details:
            runs.append((" · " + " · ".join(details), False))
        self._paragraph(runs, self.list_style, LIST_INDENT_TWIPS * (depth + 1))

    def _relationships(self, company: ExportCompany) -> None:
        self.doc.add_heading(HEADING_RELATIONSHIPS, level=2)
        rows: list[tuple[str, str, str]] = []
        for concept in self.s.concepts:
            if (
                concept.company_id == company.id
                and concept.birth == "birth"
                and concept.parent_id in self.concepts
            ):
                parent = self.concepts[concept.parent_id]
                pair = (concept, parent) if concept.birth_reverse else (parent, concept)
                action = concept.birth_action or ""
                rows.append((self._name(pair[0], concept), action, self._name(pair[1], concept)))
        for relation in self._relations(company, ("rel", "isa")):
            a, b = self.concepts[relation.a_id], self.concepts[relation.b_id]
            rows.append(
                (self._name(a, None, company), relation.action, self._name(b, None, company))
            )
        if not rows:
            self._paragraph([(NONE_TEXT, False)])
            return
        table = self._table(("Subject", "Action", "Object"))
        for row in rows:
            self._row(table, row)

    def _equivalences(self, company: ExportCompany) -> None:
        self.doc.add_heading(HEADING_EQUIVALENCES, level=2)
        relations = self._relations(company, ("same",))
        if not relations:
            self._paragraph([(NONE_TEXT, False)])
            return
        for relation in relations:
            a, b = self.concepts[relation.a_id], self.concepts[relation.b_id]
            if a.company_id != company.id:
                a, b = b, a
            other = self._name(b, None, company)
            self._paragraph(
                [(a.label, True), (f" is equivalent to {other}", False)], self.list_style
            )

    def _paragraph(
        self, runs: list[Run], style_id: str | None = None, indent: int | None = None
    ) -> None:
        """Append one paragraph of plain runs before the section properties."""
        paragraph = _paragraph_element(runs, style_id, indent)
        if self.sect_pr is not None:
            self.sect_pr.addprevious(paragraph)
        else:
            self.doc.element.body.append(paragraph)

    def _table(self, headings: tuple[str, str, str]) -> Table:
        table = self.doc.add_table(rows=0, cols=len(headings))
        table.style = "Table Grid"
        self._row(table, headings, bold=True)
        return table

    def _row(self, table: Table, values: tuple[str, ...], bold: bool = False) -> list[_Element]:
        """Append one row of text cells; returns its cells."""
        row = SubElement(table._tbl, qn("w:tr"))
        cells = []
        for value in values:
            cell = SubElement(row, qn("w:tc"))
            cell.append(_paragraph_element([(value, bold)], None, None))
            cells.append(cell)
        return cells

    def _relations(self, company: ExportCompany, kinds: tuple[str, ...]) -> list[ExportRelation]:
        return [
            r
            for r in self.s.relations
            if r.kind in kinds
            and r.a_id in self.concepts
            and r.b_id in self.concepts
            and company.id in (self.concepts[r.a_id].company_id, self.concepts[r.b_id].company_id)
        ]

    def _is_root(self, concept_id: uuid.UUID) -> bool:
        concept = self.concepts.get(concept_id)
        return concept is not None and concept.birth == "root"

    def _name(
        self,
        concept: ExportConcept,
        near: ExportConcept | None,
        company: ExportCompany | None = None,
    ) -> str:
        """The label, followed by the company's name when it is another company's concept."""
        home = near.company_id if near is not None else company.id if company else None
        if home is not None and concept.company_id != home:
            return f"{concept.label} ({self.company_names.get(concept.company_id, '')})"
        return concept.label


def _paragraph_element(runs: list[Run], style_id: str | None, indent: int | None) -> _Element:
    """A `w:p` of plain text runs, with a paragraph style and a left indent when given."""
    paragraph = OxmlElement("w:p")
    if style_id is not None or indent is not None:
        properties = SubElement(paragraph, qn("w:pPr"))
        if style_id is not None:
            SubElement(properties, qn("w:pStyle")).set(qn("w:val"), style_id)
        if indent is not None:
            SubElement(properties, qn("w:ind")).set(qn("w:left"), str(indent))
    for text, bold in runs:
        run = SubElement(paragraph, qn("w:r"))
        if bold:
            SubElement(SubElement(run, qn("w:rPr")), qn("w:b"))
        node = SubElement(run, qn("w:t"))
        node.set(qn("xml:space"), "preserve")
        node.text = _text(text)
    return paragraph


def _shade(cell: _Element, color: str) -> None:
    """Fill the cell with the colour, as a colour swatch."""
    if not _HEX.match(color):
        return
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color[1:].upper())
    properties = cell.find(qn("w:tcPr"))
    if properties is None:
        properties = OxmlElement("w:tcPr")
        cell.insert(0, properties)
    properties.append(shading)


def _text(value: str) -> str:
    return _NOT_XML.sub("", value)
